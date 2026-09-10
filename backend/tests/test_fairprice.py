"""
Fair price, 2026-09-10 (night) — the cost MODEL (backend/fairprice.py) and its seeded coefficients.

What is asserted: the seeded block in factors.json (every coefficient with a source line
beside it, the three assumptions labelled as such); the vehicle class rule; litres per
100 km laden / empty / winter; the trip formula on the fixture geometry (EUR 113.78 —
the figure in the research note) and its components; the IRU sanity band; None — never
a guess — for an unbaked route, a missing index or a blank coefficient; the fallback from
a live document without the block to the file's; the fair figure on days, weeks, totals
and Account rows; config validation; the XLSX column; and that the supplier's PDF does
NOT carry it.

WHAT THIS DOES NOT PROVE
------------------------
  * That the coefficients are RIGHT. They are public benchmarks read on 10 Sep 2026 and
    three are assumptions; the file says which. The assertions pin what was seeded so a
    silent change is caught, not that the number is true.
  * The HTTP layer is stubbed; nothing in a browser (the JS harnesses cover the section).

Run:  python3 backend/tests/test_fairprice.py
"""
import datetime
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import types
import urllib.error

HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
ROOT = os.path.dirname(BACKEND)
sys.path.insert(0, BACKEND)

# ---- stubs, verbatim from test_lookahead.py --------------------------------------
_fp = types.ModuleType("flexpolyline")
_fp.decode = lambda s: []
_fp.encode = lambda pts: ""
sys.modules.setdefault("flexpolyline", _fp)


def _passthrough_decorator(*a, **k):
    def wrap(fn):
        return fn
    return wrap


class _App:
    def __init__(self, *a, **k):
        pass

    def get(self, *a, **k):
        return _passthrough_decorator()

    post = put = delete = patch = get

    def add_middleware(self, *a, **k):
        pass

    def mount(self, *a, **k):
        pass


class _HTTPException(Exception):
    def __init__(self, status_code, detail=""):
        super().__init__(f"{status_code}: {detail}")
        self.status_code = status_code
        self.detail = detail


def _Query(default=None, **k):
    return default


_fa = types.ModuleType("fastapi")
_fa.FastAPI = _App
_fa.HTTPException = _HTTPException
_fa.Query = _Query
sys.modules.setdefault("fastapi", _fa)
_mw = types.ModuleType("fastapi.middleware")
_cors = types.ModuleType("fastapi.middleware.cors")
_cors.CORSMiddleware = object
_mw.cors = _cors
sys.modules.setdefault("fastapi.middleware", _mw)
sys.modules.setdefault("fastapi.middleware.cors", _cors)
_resp = types.ModuleType("fastapi.responses")
_resp.FileResponse = lambda *a, **k: None
sys.modules.setdefault("fastapi.responses", _resp)
_static = types.ModuleType("fastapi.staticfiles")


class _StaticFiles:
    def __init__(self, *a, **k):
        pass

    async def get_response(self, path, scope):
        return None


_static.StaticFiles = _StaticFiles
sys.modules.setdefault("fastapi.staticfiles", _static)

TMP = tempfile.mkdtemp(prefix="rbe_fairprice_")
os.environ.pop("DATABASE_URL", None)
os.environ.pop("ADMIN_TOKEN", None)

import db  # noqa: E402
db._SQLITE_PATH = os.path.join(TMP, "scratch.db")
import conversions  # noqa: E402
import network  # noqa: E402
import weeks  # noqa: E402
import main  # noqa: E402
import access  # noqa: E402
import days  # noqa: E402
import derived  # noqa: E402
import costing  # noqa: E402
import fuel  # noqa: E402
import lookahead  # noqa: E402
import export  # noqa: E402
import fairprice  # noqa: E402
import config  # noqa: E402
for _v in ("IPT1_CODE", "IPT2_CODE", "IPT3_CODE", "IPT4_CODE", "IPT5_CODE", "IPT6_CODE",
           "PLANNER_CODE", "ADMIN_CODE"):
    os.environ.pop(_v, None)
access.set_current("planner123")
days.START_YEAR = main.START_YEAR

PASS = 0
FAIL = []


def ok(label, cond, extra=""):
    global PASS
    if cond:
        PASS += 1
    else:
        FAIL.append(f"{label} {extra}".strip())


def _raises(fn):
    try:
        fn()
    except Exception as e:
        return e
    return None


def reset_db():
    if os.path.exists(db._SQLITE_PATH):
        os.remove(db._SQLITE_PATH)
    db.init_db()
    db.init_network_db()
    db.init_taxonomy_db()
    db.init_zones_db()
    db.init_gates_db()
    db.init_weeks_db()
    db.init_lookahead_db()
    db.init_config_db()
    db.init_costing_db()
    db.init_tenant()
    import config as _cfg
    _cfg.invalidate()
    costing.invalidate()




def _code_tokens(src):
    """NAME and NUMBER tokens only — docstrings and comments may NAME a thing the code must not hold."""
    import io as _io
    import tokenize
    names, nums = set(), set()
    for t in tokenize.generate_tokens(_io.StringIO(src).readline):
        if t.type == tokenize.NAME:
            names.add(t.string.lower())
        elif t.type == tokenize.NUMBER:
            try:
                nums.add(float(t.string))
            except ValueError:
                pass
    return names, nums


FACT_FILE = json.load(open(os.path.join(BACKEND, "factors.json"), encoding="utf-8"))
FP = FACT_FILE.get("fair_price") or {}

# =========================================================================== #
#  0. The seeded block — sourced, labelled, and pinned                         #
# =========================================================================== #
ok("factors.json carries a fair_price block with a README", isinstance(FP, dict) and len(FP.get("_README", "")) > 100)
ok("...every coefficient has a source line beside it",
   "_source" in (FP.get("consumption") or {}) and "_driver_source" in FP and "_vehicle_standing_source" in FP
   and "_running_source" in FP and "_margin_source" in FP and "_source" in (FP.get("season") or {}))
ok("🔴 the three unsourced numbers are LABELLED as assumptions in the file",
   all("ASSUMPTION" in FP.get(k, "") for k in ("_vehicle_standing_source", "_running_source", "_margin_source"))
   and "ASSUMPTION" not in FP.get("_driver_source", ""))
ok("...and the sourced ones name where they came from",
   "ICCT" in FP["consumption"]["_source"] and "DfT" in FP["consumption"]["_source"]
   and "Statistics Estonia" in FP["_driver_source"] and "Transpordiamet" in FP["season"]["_source"]
   and "fueleconomy.gov" in FP["season"]["_source"])
ok("the seeded values are what the research note said (a silent change is caught here)",
   FP["driver_eur_per_h"] == 20.0 and FP["vehicle_standing_eur_per_h"] == 16.0 and FP["running_eur_per_km"] == 0.12
   and FP["margin_pct"] == 8.0 and FP["consumption"]["l_per_100km_per_tonne"] == 0.4
   and FP["consumption"]["rigid"]["l_per_100km_empty"] == 27.0 and FP["consumption"]["artic"]["l_per_100km_empty"] == 23.6
   and FP["season"]["winter_months"] == [11, 12, 1, 2, 3] and FP["season"]["winter_consumption_uplift_pct"] == 8.0
   and FP["season"]["thaw_months"] == [3, 4])
ok("the standing-cost derivation is written down and arithmetically consistent (190k, 7 yr, 20 % residual, 3 %, vignette 1,100, 1,800 h)",
   abs((190000 * 0.8 / 7 + 190000 * 0.03 + 1100) / 1800 - 15.8) < 0.1 and FP.get("vehicle_new_price_eur") == 190000)
ok("🔴 the diesel price is NOT in the block — it is the live index",
   not any(("diesel" in k.lower() or "eur_per_l" in k.lower()) for k in FP if not k.startswith("_")))
_fsrc = open(os.path.join(BACKEND, "fairprice.py"), encoding="utf-8").read()
ok("fairprice.py reads no database — pure arithmetic on what derived.py hands it",
   "import db" not in _fsrc and "db." not in _fsrc.replace("db.py", ""))

# =========================================================================== #
#  1. The pure functions                                                       #
# =========================================================================== #
P = fairprice.params(FACT_FILE)
ok("params() strips the _source lines and keeps the numbers",
   P["driver_eur_per_h"] == 20.0 and "_README" not in P and "_source" not in P["consumption"]
   and P["consumption"]["rigid"]["l_per_100km_empty"] == 27.0)
ok("params() on a document WITHOUT the block falls back to the file's, key by key",
   fairprice.params({"vehicles": {}})["driver_eur_per_h"] == 20.0
   and fairprice.params({"fair_price": {"driver_eur_per_h": 25}})["driver_eur_per_h"] == 25
   and fairprice.params({"fair_price": {"driver_eur_per_h": 25}})["running_eur_per_km"] == 0.12
   and fairprice.params({"fair_price": {"consumption": {"artic": {"l_per_100km_empty": 22}}}})["consumption"]["rigid"]["l_per_100km_empty"] == 27.0)
V8 = "Rigid 8-wheeler (32t)"
V12 = "N3 tractor (BC) + O4 tipping semi (DA) GCW 40/44 t"
ok("vehicle_class: a rigid is 'rigid'; a tractor + O4 semi, a 44 t artic, are 'artic'; unknown ⇒ rigid",
   fairprice.vehicle_class(V8, FACT_FILE) == "rigid" and fairprice.vehicle_class(V12, FACT_FILE) == "artic"
   and fairprice.vehicle_class("Artic Tipper (44t)", FACT_FILE) == "artic" and fairprice.vehicle_class("Rigid 7.5t", FACT_FILE) == "rigid"
   and fairprice.vehicle_class("Ghost truck", FACT_FILE) == "rigid")
ok("litres: rigid empty 27, with 20 t 35 (= 27 + 0.4 × 20), artic with 26 t 34.0 (= the ICCT figure)",
   fairprice.litres_per_100km("rigid", 0, P) == 27.0 and abs(fairprice.litres_per_100km("rigid", 20, P) - 35.0) < 1e-9
   and abs(fairprice.litres_per_100km("artic", 26, P) - 34.0) < 1e-9)
ok("...winter: January × 1.08; September not", abs(fairprice.litres_per_100km("rigid", 20, P, month=1) - 37.8) < 1e-9
   and abs(fairprice.litres_per_100km("rigid", 20, P, month=9) - 35.0) < 1e-9)
ok("...a blank coefficient ⇒ None, never a guess",
   fairprice.litres_per_100km("rigid", 20, {"consumption": {"l_per_100km_per_tonne": None, "rigid": {"l_per_100km_empty": 27}}}) is None)
CTX = {"baked": True, "distance_km": 30.0, "return_km": 30.0, "cycle_min": 104.0, "payload_t": 20.0, "vehicle_class": "rigid"}
t = fairprice.trip_cost(CTX, 9, 1.922, P)
ok("⭐ the fixture trip (60 km round trip, 20 t, 1.73 h, diesel 1.922) costs € 113.78 — the research note's figure",
   t is not None and t["eur_trip"] == 113.78, str(t))
ok("...its components: 18.6 L → € 35.75 fuel · € 62.40 time · € 7.20 running · € 8.43 margin; fuel 33.9 % of cost",
   t and t["litres_trip"] == 18.6 and t["fuel_eur"] == 35.75 and t["time_eur"] == 62.4 and t["running_eur"] == 7.2
   and t["margin_eur"] == 8.43 and t["fuel_share_pct"] == 33.9)
ok("...€ 5.69 / t and € 1.896 / km — inside the IRU 0.50–2.00 €/km band", t and t["eur_per_t"] == 5.69 and 0.5 <= t["eur_per_km"] <= 2.0)
ok("...the same arithmetic by hand: (fuel + time + running) × 1.08",
   t and abs((t["fuel_eur"] + t["time_eur"] + t["running_eur"]) * 1.08 - t["eur_trip"]) < 0.02)
tw = fairprice.trip_cost(CTX, 1, 1.922, P)
ok("January costs more than September (the uplift on the fuel term only) and carries the WINTER flag",
   tw["eur_trip"] > t["eur_trip"] and abs((tw["eur_trip"] - t["eur_trip"]) - t["fuel_eur"] * 0.08 * 1.08) < 0.02
   and tw["flags"] == ["WINTER"] and t["flags"] == [])
ok("March carries BOTH flags (winter month and thaw month); April thaw only",
   fairprice.trip_cost(CTX, 3, 1.922, P)["flags"] == ["WINTER", "THAW"] and fairprice.trip_cost(CTX, 4, 1.922, P)["flags"] == ["THAW"])
ta = fairprice.trip_cost({**CTX, "vehicle_class": "artic", "payload_t": 26.0}, 9, 1.922, P)
ok("an artic carrying 26 t: 34.0 laden / 23.6 empty, cheaper per tonne than the rigid", ta["l_per_100km_laden"] == 34.0
   and ta["l_per_100km_empty"] == 23.6 and ta["eur_per_t"] < t["eur_per_t"])
ok("🔴 unbaked ⇒ None; no index ⇒ None; a blank coefficient ⇒ None — nothing invented",
   fairprice.trip_cost({**CTX, "baked": False}, 9, 1.922, P) is None and fairprice.trip_cost(CTX, 9, None, P) is None
   and fairprice.trip_cost(CTX, 9, 1.922, {**P, "driver_eur_per_h": None}) is None)
tr = fairprice.trip_cost({**CTX, "return_km": None}, 9, 1.922, P)
ok("a return leg that is not baked: the outbound stands in for its distance (as the cycle already does)",
   tr is not None and tr["litres_trip"] == t["litres_trip"])
ok("a higher diesel index raises only the fuel term", fairprice.trip_cost(CTX, 9, 2.922, P)["time_eur"] == t["time_eur"]
   and fairprice.trip_cost(CTX, 9, 2.922, P)["fuel_eur"] > t["fuel_eur"])
S = fairprice.summary(P, 1.922)
ok("summary() reports the coefficients in force and 'complete'", S["complete"] is True and S["driver_eur_per_h"] == 20.0
   and fairprice.summary(P, None)["complete"] is False)

# =========================================================================== #
#  2. On the Look-ahead — days, weeks, totals, Account                         #
# =========================================================================== #
reset_db()
access.set_current("planner123")
db.execute("INSERT INTO locations (id, name, lat, lon) VALUES (?, ?, ?, ?)", ("L1", "Pit", 58.5, 24.0))
db.execute("INSERT INTO locations (id, name, lat, lon) VALUES (?, ?, ?, ?)", ("L2", "Site", 58.6, 24.4))
for _rid in ("R1", "R2", "R3"):
    db.execute("INSERT INTO routes (id, origin_id, dest_id, ipt) VALUES (?, ?, ?, ?)", (_rid, "L1", "L2", "IPT 1"))


def _geom(rid, prof, leg, km, hr):
    db.execute("INSERT INTO route_geometry (tenant_id, route_id, vehicle_profile, leg, alt_index, "
               "geometry, distance_km, duration_hr) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
               ("default", rid, prof, leg, 0, "[[24,58.5],[24.4,58.6]]", km, hr))


_geom("R1", V8, "loaded", 30.0, 0.75); _geom("R1", V8, "return", 30.0, 0.65)
_geom("R2", V12, "loaded", 30.0, 0.75); _geom("R2", V12, "return", 30.0, 0.65)
MONTHS = tuple(range(1, 25))


def _line(rid, sect, qty, veh):
    main.save_matrix_row(main.MatrixRow(
        route_id=rid, discipline="earthworks", section_id=sect, material_type="Small aggregate",
        material_description=None, vehicle_type=veh, submitted_by="tester", unit="t",
        status="Pending", cells=[main.Cell(month_index=m, quantity=qty) for m in MONTHS], ipt="IPT1"))
    main.set_route_status(rid, main.StatusUpdate(status="Approved"), discipline="earthworks", section_id=sect)


_line("R1", "WS1", 4000.0, V8)       # rigid, baked, no rate
_line("R2", "WS2", 4000.0, V12)      # artic, baked, no rate
_line("R3", "WS3", 900.0, V8)        # unbaked
res = main.list_forecast_days()
L = {l["section_id"]: l for l in res["lines"]}
ok("🔴 with NO index row there is no fair price anywhere — the model needs the live diesel price",
   all(d["derived"]["fair_eur"] is None for l in res["lines"] for d in l["days"]) and res["totals"]["fair_eur"] is None
   and res["costing"]["fair"]["complete"] is False and L["WS1"]["context"]["fair"] is None)
fuel.set_manual(1.922, "2026-09-07", by="admin")
res = main.list_forecast_days()
L = {l["section_id"]: l for l in res["lines"]}
c1 = L["WS1"]["context"]
ok("⭐ with the index: R1 (rigid, 20 t, 60 km, cycle from HERE) carries a fair price per trip, class 'rigid', its calendar month",
   c1["fair"] is not None and c1["vehicle_class"] == "rigid" and c1["fair"]["eur_trip"] > 0
   and c1["month"] == ((L["WS1"]["month_index"] - 1) % 12) + 1, str(c1.get("fair")))
d1 = [d for d in L["WS1"]["days"] if datetime.date.fromisoformat(d["day_date"]).weekday() <= 4][0]["derived"]
ok("...a day's fair € = trips × € / trip — the same trips slice 2 derives",
   abs(d1["fair_eur"] - round(d1["trips"] * c1["fair"]["eur_trip"], 2)) < 0.011)
ok("...no quote, no target — fair is the ONLY figure on the line, and it is its own field",
   d1["eur"] is None and d1["eur_adj"] is None and d1["fair_eur"] is not None and c1["rate_source"] is None)
ok("...the cycle the model used is the line's own (context.cycle_min)", c1["fair"]["month"] == c1["month"]
   and abs(c1["fair"]["time_eur"] - round(c1["cycle_min"] / 60.0 * (20.0 + 16.0), 2)) < 0.011)
ok("R2 (artic, 26 t) prices as an artic", L["WS2"]["context"]["vehicle_class"] == "artic"
   and L["WS2"]["context"]["fair"]["l_per_100km_laden"] == 34.0)
ok("🔴 R3 (unbaked) has NO fair price — no distance, nothing invented", L["WS3"]["context"]["fair"] is None
   and all(d["derived"]["fair_eur"] is None for d in L["WS3"]["days"]) and L["WS3"]["week_derived"]["fair_eur"] is None)
ok("week and totals sum the fair figure over the lines that have one: 2 of 3",
   L["WS1"]["week_derived"]["fair_eur"] is not None and res["totals"]["fair_lines"] == 2
   and abs(res["totals"]["fair_eur"] - (L["WS1"]["week_derived"]["fair_eur"] + L["WS2"]["week_derived"]["fair_eur"])) < 0.05)
ok("...and a typed target does not touch the fair figure (three independent columns)",
   (costing.set_target({"rate_eur_per_t": 3.0}, by="a") or True)
   and {l["section_id"]: l for l in main.list_forecast_days()["lines"]}["WS1"]["week_derived"]["fair_eur"] == L["WS1"]["week_derived"]["fair_eur"]
   and {l["section_id"]: l for l in main.list_forecast_days()["lines"]}["WS1"]["week_derived"]["eur"] is not None)
pg = lookahead.page(bucket="commit")
if pg["account"].get("rows"):
    A = {r["section_id"]: r for r in pg["account"]["rows"]}
    ok("account rows carry planned_fair_eur and the season flags; the unbaked row None",
       A["WS1"]["planned_fair_eur"] is not None and isinstance(A["WS1"]["fair_flags"], list) and A["WS3"]["planned_fair_eur"] is None
       and pg["account"]["fair_lines"] == 2, str({k: v.get("planned_fair_eur") for k, v in A.items()}))
else:
    ok("(no account week in this bucket — the account assertion did not run)", False)
ok("the page's costing block carries the model's coefficients in force",
   pg["costing"]["fair"]["driver_eur_per_h"] == 20.0 and pg["costing"]["fair"]["complete"] is True)
# the live document without the block prices from the file; with an override it prices from the override
config.save({**FACT_FILE, "fair_price": {**FP, "driver_eur_per_h": 40.0}}, by="t")
res2 = main.list_forecast_days()
c1b = {l["section_id"]: l for l in res2["lines"]}["WS1"]["context"]["fair"]
ok("⭐ a coefficient changed on Config moves the figure on the next read (driver 20 → 40: time term doubles the driver share)",
   c1b is not None and abs(c1b.get("time_eur", 0) - round(c1["cycle_min"] / 60.0 * (40.0 + 16.0), 2)) < 0.011
   and c1b.get("eur_trip", 0) > c1["fair"]["eur_trip"])
_doc_no_block = {k: v for k, v in FACT_FILE.items() if k != "fair_price"}
config.save(_doc_no_block, by="t")
c1c = {l["section_id"]: l for l in main.list_forecast_days()["lines"]}["WS1"]["context"]["fair"]
ok("...and a live document WITHOUT the block prices from the file's seeded values", (c1c or {}).get("eur_trip") == c1["fair"]["eur_trip"])
config.reset_to_file(conversions, by="t")

# =========================================================================== #
#  3. Validation, export, what must not exist                                  #
# =========================================================================== #
ok("config.validate: a negative driver rate, a month 13, a zero empty consumption are refused; the block absent passes",
   any("driver_eur_per_h" in x for x in config.validate({**FACT_FILE, "fair_price": {**FP, "driver_eur_per_h": -1}}))
   and any("winter_months" in x for x in config.validate({**FACT_FILE, "fair_price": {**FP, "season": {**FP["season"], "winter_months": [13]}}}))
   and any("l_per_100km_empty" in x for x in config.validate({**FACT_FILE, "fair_price": {**FP, "consumption": {**FP["consumption"], "rigid": {"l_per_100km_empty": 0}}}}))
   and not config.validate(_doc_no_block) and not config.validate(FACT_FILE))
pg = lookahead.page(bucket="commit")
xb = export.build_xlsx(pg)
import openpyxl as _ox
wb = _ox.load_workbook(io.BytesIO(xb))
ws = wb["Commit week"]
hdr = [ws.cell(row=1, column=j).value for j in range(1, ws.max_column + 1)]
ok("XLSX: '€ fair (model)', 'fair €/t' and 'fair flags' columns follow '€ + BAF'",
   "€ fair (model)" in hdr and hdr.index("€ fair (model)") == hdr.index("€ + BAF") + 1 and "fair €/t" in hdr and "fair flags" in hdr)
cf = hdr.index("€ fair (model)") + 1
vals = [ws.cell(row=i, column=cf).value for i in range(2, ws.max_row + 1)]
ok("...baked rows carry a fair figure and the unbaked rows are blank", sum(1 for v in vals if v is not None) == 10 and sum(1 for v in vals if v is None) == 5)
wf = wb["Fuel"]
ftxt = " ".join(str(wf.cell(row=i, column=j).value or "") for i in range(1, wf.max_row + 1) for j in (1, 2))
ok("...the Fuel sheet states the coefficients the fair figures used", "Fair price: driver" in ftxt and "Fair price: margin" in ftxt and "winter months" in ftxt)
pb = export.build_pdf(pg)
_txt = subprocess.run(["pdftotext", "-layout", "-", "-"], input=pb, capture_output=True).stdout.decode("utf-8") if shutil.which("pdftotext") else pb.decode("latin-1")
ok("🔴 the supplier's PDF does NOT carry the fair price — it is the planner's negotiating figure",
   "fair" not in _txt.lower() and "model" not in _txt.lower())
_fe = open(os.path.join(ROOT, "frontend", "index.html"), encoding="utf-8").read()
ok("the page: FairPriceSection defined once and mounted in the Costing tab; the Account column and the KPI caption say 'model'",
   len(re.findall(r"function FairPriceSection\(", _fe)) == 1 and _fe.count("<FairPriceSection") == 1
   and "Fair € <span" in _fe and "(model, ${totals.fair_lines} of ${totals.lines})" in _fe)
ok("...the copy calls it a floor for negotiation, never a quote", "a floor for negotiation, never a quote" in _fe)
ok("🔴 no fair price on the public map", "fair price" not in open(os.path.join(ROOT, "map", "index.html"), encoding="utf-8").read().lower())
ok("no route type field was invented: no 'route_type' / 'road_surface' NAME in the backend (docstrings may say the word)",
   all(not any(x in n for n in _code_tokens(open(os.path.join(BACKEND, f), encoding="utf-8").read())[0] for x in ("route_type", "road_surface", "surface_factor"))
       for f in ("fairprice.py", "derived.py", "lookahead.py")))

print()
for f in FAIL:
    print("  FAIL:", f)
print(f"\n{PASS} passed, {len(FAIL)} failed")
sys.exit(1 if FAIL else 0)
