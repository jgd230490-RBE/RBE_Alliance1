"""
Costing, 2026-09-10 (evening) — target rates and the fuel index + BAF.

What is asserted: the feed parser on a fixture of the real EuroOilWatch payload; a
refresh against a STUBBED urlopen (success, 500, timeout, garbage), the last good row
surviving a failure, empty-cache-plus-failure reading as "unavailable" not 500; the
manual index; the tenant settings and their validation; the yard price never touching
the bulletin row; BAF null without a share or a base; Confirm week locking the base
ONCE; the target rate pricing an unpriced route and never mixing with a typed one; the
Quote + BAF figure beside the quote on days, weeks, totals and the account rows; the
export columns and the Fuel sheet without the planner's own numbers; and the things the
note says must not exist (fuel on the public map, an upload endpoint, a seeded price).

WHAT THIS DOES NOT PROVE
------------------------
  * eurooilwatch.com is never called. The parser runs on a fixture captured on 10 Sep
    2026 (EE diesel 1.922, bulletin 2026-09-07); the fetch is exercised only against a
    stub. Whether Render can reach the feed is checked on the deployment.
  * The HTTP layer is stubbed (same harness as test_lookahead.py): endpoint BODIES run
    as a planner; nothing proves a route is mounted or a query string is parsed.
  * No Postgres branch. fuel_index is created on SQLite only.
  * Nothing in a browser — parse_frontend.js / render_frontend.js cover the widget.

Run:  python3 backend/tests/test_costing.py
"""
import datetime
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import types
import urllib.error

HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
ROOT = os.path.dirname(BACKEND)
sys.path.insert(0, BACKEND)

# ---- stubs, verbatim from test_lookahead.py --------------------------------------
_fp = types.ModuleType("flexpolyline")
_fp.decode = lambda s: []
_fp.encode = lambda pts: ""
sys.modules.setdefault("flexpolyline", _fp)


def _passthrough_decorator(*a, **k):
    def wrap(fn):
        return fn
    return wrap


class _App:
    def __init__(self, *a, **k):
        pass

    def get(self, *a, **k):
        return _passthrough_decorator()

    post = put = delete = patch = get

    def add_middleware(self, *a, **k):
        pass

    def mount(self, *a, **k):
        pass


class _HTTPException(Exception):
    def __init__(self, status_code, detail=""):
        super().__init__(f"{status_code}: {detail}")
        self.status_code = status_code
        self.detail = detail


def _Query(default=None, **k):
    return default


_fa = types.ModuleType("fastapi")
_fa.FastAPI = _App
_fa.HTTPException = _HTTPException
_fa.Query = _Query
sys.modules.setdefault("fastapi", _fa)
_mw = types.ModuleType("fastapi.middleware")
_cors = types.ModuleType("fastapi.middleware.cors")
_cors.CORSMiddleware = object
_mw.cors = _cors
sys.modules.setdefault("fastapi.middleware", _mw)
sys.modules.setdefault("fastapi.middleware.cors", _cors)
_resp = types.ModuleType("fastapi.responses")
_resp.FileResponse = lambda *a, **k: None
sys.modules.setdefault("fastapi.responses", _resp)
_static = types.ModuleType("fastapi.staticfiles")


class _StaticFiles:
    def __init__(self, *a, **k):
        pass

    async def get_response(self, path, scope):
        return None


_static.StaticFiles = _StaticFiles
sys.modules.setdefault("fastapi.staticfiles", _static)

TMP = tempfile.mkdtemp(prefix="rbe_costing_")
os.environ.pop("DATABASE_URL", None)
os.environ.pop("ADMIN_TOKEN", None)

import db  # noqa: E402
db._SQLITE_PATH = os.path.join(TMP, "scratch.db")
import conversions  # noqa: E402
import network  # noqa: E402
import weeks  # noqa: E402
import main  # noqa: E402
import access  # noqa: E402
import days  # noqa: E402
import derived  # noqa: E402
import costing  # noqa: E402
import fuel  # noqa: E402
import lookahead  # noqa: E402
import export  # noqa: E402
for _v in ("IPT1_CODE", "IPT2_CODE", "IPT3_CODE", "IPT4_CODE", "IPT5_CODE", "IPT6_CODE",
           "PLANNER_CODE", "ADMIN_CODE"):
    os.environ.pop(_v, None)
access.set_current("planner123")
days.START_YEAR = main.START_YEAR

PASS = 0
FAIL = []


def ok(label, cond, extra=""):
    global PASS
    if cond:
        PASS += 1
    else:
        FAIL.append(f"{label} {extra}".strip())


def _raises(fn):
    try:
        fn()
    except Exception as e:
        return e
    return None


def reset_db():
    if os.path.exists(db._SQLITE_PATH):
        os.remove(db._SQLITE_PATH)
    db.init_db()
    db.init_network_db()
    db.init_taxonomy_db()
    db.init_zones_db()
    db.init_gates_db()
    db.init_weeks_db()
    db.init_lookahead_db()
    db.init_config_db()
    db.init_costing_db()
    db.init_tenant()
    import config as _cfg
    _cfg.invalidate()
    costing.invalidate()


# The real payload shape, captured from https://eurooilwatch.com/api/v1/prices on
# 2026-09-10 (bulletinDate and dataSource are TOP-LEVEL; the note implied per country).
FEED = {"lastUpdated": "2026-09-10T09:37:05.244Z", "bulletinDate": "2026-09-07",
        "dataSource": "EC Weekly Oil Bulletin (2026-09-07)",
        "countries": [
            {"countryCode": "LV", "countryName": "Latvia", "petrolPrice": 1.70, "dieselPrice": 1.66,
             "petrolChangePct": None, "dieselChangePct": None},
            {"countryCode": "EE", "countryName": "Estonia", "petrolPrice": 1.812, "dieselPrice": 1.922,
             "petrolChangePct": None, "dieselChangePct": None}],
        "euAverage": {"petrolPrice": 1.65, "dieselPrice": 1.58}}

# =========================================================================== #
#  0. Registration — global by decision, and declared as such                  #
# =========================================================================== #
reset_db()
ok("fuel_index exists after init_costing_db()",
   any(r["name"] == "fuel_index" for r in db.query("SELECT name FROM sqlite_master WHERE type='table'")))
ok("🔴 fuel_index is NOT tenanted — a national index is not client data",
   "fuel_index" not in db.TENANTED_TABLES and "fuel_index" not in db._TENANT_DDL
   and "tenant_id" not in [c.lower() for c in db._columns_of(db.get_conn().cursor(), "fuel_index")])
_audit = open(os.path.join(HERE, "test_tenant_audit.py"), encoding="utf-8").read()
ok("...and the audit lists it as untenanted WITH a reason (not silently exempt)",
   re.search(r'"fuel_index":\s*"[^"]{20,}"', _audit) is not None)
ok("the tenant's settings live in the tenanted config table under one key, no new table",
   costing.KEY == "costing" and len(db.TENANTED_TABLES) == 16)
ok("nothing is seeded: no target rate, no share, no base, no yard, no index row",
   costing.settings()["target"] == {k: None for k in costing.TARGET_FIELDS}
   and costing.settings()["fuel"]["share_pct"] is None and costing.settings()["fuel"]["baf_base_eur_per_l"] is None
   and costing.settings()["fuel"]["yard_eur_per_l"] is None and fuel.get_index("EE") is None)
_src = {f: open(os.path.join(BACKEND, f), encoding="utf-8").read() for f in ("fuel.py", "costing.py", "derived.py", "main.py", "weeks.py")}


def _code_tokens(src):
    """NAME and NUMBER tokens only — docstrings and comments may NAME a thing the code must not hold."""
    import io as _io
    import tokenize
    names, nums = set(), set()
    for t in tokenize.generate_tokens(_io.StringIO(src).readline):
        if t.type == tokenize.NAME:
            names.add(t.string.lower())
        elif t.type == tokenize.NUMBER:
            try:
                nums.add(float(t.string))
            except ValueError:
                pass
    return names, nums


_tok = {f: _code_tokens(_src[f]) for f in _src}
ok("🔴 no litre price and no share is written into the CODE as a default (no 1.9x, no 25 — only the docstrings name them)",
   not any(1.85 <= n <= 2.1 and n != int(n) for n in _tok["costing.py"][1] | _tok["fuel.py"][1])
   and 25.0 not in _tok["costing.py"][1] and 25.0 not in _tok["fuel.py"][1])

# =========================================================================== #
#  1. The parser, on the captured payload                                      #
# =========================================================================== #
p = fuel.parse_feed(FEED, "EE")
ok("⭐ EE diesel parsed from the fixture: 1.922 on bulletin 2026-09-07",
   p.get("eur_per_l") == 1.922 and p.get("bulletin_date") == "2026-09-07", str(p))
ok("...and the source label is the Commission bulletin the feed names",
   p.get("raw_source_label") == "EC Weekly Oil Bulletin (2026-09-07)")
ok("another country parses from the same document", fuel.parse_feed(FEED, "LV").get("eur_per_l") == 1.66)
ok("a country the feed lacks is an error, not a zero", "error" in fuel.parse_feed(FEED, "XX"))
_bad = json.loads(json.dumps(FEED)); _bad["countries"][1]["dieselPrice"] = None
ok("a missing diesel price is an error, not a zero", "error" in fuel.parse_feed(_bad, "EE"))
_bad = json.loads(json.dumps(FEED)); _bad["bulletinDate"] = "soon"
ok("a bulletin date that is not a date is an error", "error" in fuel.parse_feed(_bad, "EE"))
ok("a non-object payload is an error", "error" in fuel.parse_feed("<html>", "EE") and "error" in fuel.parse_feed(None, "EE"))
ok("the lower-case country code still matches", fuel.parse_feed(FEED, "ee").get("eur_per_l") == 1.922)

# =========================================================================== #
#  2. Refresh against a stubbed feed — success, failure, empty cache           #
# =========================================================================== #
_calls = {"n": 0, "mode": "ok", "timeout": None}


class _Resp:
    def __init__(self, body):
        self.body = body

    def read(self):
        return self.body

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def _fake_urlopen(req, timeout=None):
    _calls["n"] += 1
    _calls["timeout"] = timeout
    if _calls["mode"] == "ok":
        return _Resp(json.dumps(FEED).encode("utf-8"))
    if _calls["mode"] == "500":
        raise urllib.error.HTTPError(fuel.FEED_URL, 500, "Internal Server Error", {}, None)
    if _calls["mode"] == "timeout":
        raise TimeoutError("timed out")
    if _calls["mode"] == "garbage":
        return _Resp(b"<html>maintenance</html>")
    raise RuntimeError("unreachable")


_real_urlopen = fuel.urllib.request.urlopen
fuel.urllib.request.urlopen = _fake_urlopen
try:
    # empty cache + failure ⇒ no price, stale, an error recorded, NOT an exception
    _calls["mode"] = "500"
    st = fuel.refresh("EE", sync=True)
    row = fuel.get_index("EE")
    ok("🔴 empty cache + feed 500 ⇒ status 'unavailable', no exception",
       st.get("status") == "unavailable" and "500" in (st.get("error") or ""))
    ok("...the row records the failed ATTEMPT and no price", row is not None and row.get("eur_per_l") is None
       and row.get("last_error") and row.get("fetched_at") is None)
    S = fuel.state("EE")
    ok("...and the widget's state reads eur_per_l None + stale True — 'Index unavailable', never a guess",
       S["eur_per_l"] is None and S["stale"] is True and S["last_error"])
    ok("the fetch uses the 8 s timeout the note fixed", _calls["timeout"] == 8 and fuel.TIMEOUT_S == 8)

    # success
    _calls["mode"] = "ok"
    st = fuel.refresh("EE", sync=True)
    row = fuel.get_index("EE")
    ok("⭐ a good fetch stores the bulletin row: 1.922 / 2026-09-07 / source bulletin",
       st.get("status") == "ok" and row["eur_per_l"] == 1.922 and row["bulletin_date"] == "2026-09-07"
       and row["source"] == fuel.SOURCE_BULLETIN and row["fetched_at"] and row["last_error"] is None, str(row))
    ok("...one row per country, still ONE row after two refreshes",
       db.query("SELECT COUNT(*) AS n FROM fuel_index")[0]["n"] == 1)
    ok("...state: not stale (bulletin 2026-09-07 is within 8 days of 10 Sep), attribution names the bulletin",
       fuel.state("EE", today=datetime.date(2026, 9, 10))["stale"] is False
       and fuel.ATTRIBUTION == "EU Weekly Oil Bulletin via EuroOilWatch")
    ok("...stale once the bulletin is older than 8 days — judged on the BULLETIN, not the fetch",
       fuel.state("EE", today=datetime.date(2026, 9, 16))["stale"] is True
       and fuel.is_stale({"eur_per_l": 1.9, "bulletin_date": "2026-09-07"}, today=datetime.date(2026, 9, 15)) is False)

    # failure AFTER a good row: the good row stands
    _calls["mode"] = "timeout"
    st = fuel.refresh("EE", sync=True)
    row2 = fuel.get_index("EE")
    ok("🔴 a timeout keeps the last good row (price, date, source unchanged) and records the error",
       st.get("status") == "unavailable" and row2["eur_per_l"] == 1.922 and row2["bulletin_date"] == "2026-09-07"
       and row2["fetched_at"] == row["fetched_at"] and "Timeout" in (row2.get("last_error") or ""), str(row2))
    _calls["mode"] = "garbage"
    fuel.refresh("EE", sync=True)
    ok("...and a non-JSON body does the same — the good row is never overwritten by junk",
       fuel.get_index("EE")["eur_per_l"] == 1.922 and fuel.get_index("EE")["last_error"])

    # needs_refresh: the 12 h rule on the last ATTEMPT (a failing feed is not hammered)
    ok("needs_refresh: no row ⇒ yes; attempt 1 h ago ⇒ no; attempt 13 h ago ⇒ yes",
       fuel.needs_refresh(None) is True
       and fuel.needs_refresh({"last_attempt_at": (datetime.datetime.utcnow() - datetime.timedelta(hours=1)).isoformat() + "Z"}) is False
       and fuel.needs_refresh({"last_attempt_at": (datetime.datetime.utcnow() - datetime.timedelta(hours=13)).isoformat() + "Z"}) is True)
    _calls["n"] = 0
    fuel.ensure_fresh("EE", sync=True)
    ok("ensure_fresh does NOT call the feed when the last attempt is recent", _calls["n"] == 0)
    db.execute("UPDATE fuel_index SET last_attempt_at = ? WHERE country = ?", ("2026-09-01T00:00:00Z", "EE"))
    _calls["mode"] = "ok"
    fuel.ensure_fresh("EE", sync=True)
    ok("...and DOES when it is older than 12 h", _calls["n"] == 1)

    # manual
    r = fuel.set_manual(1.85, "2026-09-14", by="admin")
    row3 = fuel.get_index("EE")
    ok("⭐ a manual index writes the SAME row with source 'manual'",
       r["ok"] and row3["source"] == "manual" and row3["eur_per_l"] == 1.85 and row3["bulletin_date"] == "2026-09-14"
       and db.query("SELECT COUNT(*) AS n FROM fuel_index")[0]["n"] == 1)
    ok("...a bad manual price or date is refused",
       not fuel.set_manual(0, "2026-09-14")["ok"] and not fuel.set_manual(1.8, "next week")["ok"]
       and not fuel.set_manual("abc", "2026-09-14")["ok"])
    # the background path: a thread, and the state says it ran
    _calls["mode"] = "ok"
    st = fuel.refresh("EE", sync=False)
    import time as _t
    for _ in range(50):
        if not fuel.refresh_state()["running"]:
            break
        _t.sleep(0.05)
    ok("refresh(sync=False) runs in a thread and lands the bulletin row again (manual overwritten by a fresh fetch)",
       fuel.refresh_state()["running"] is False and fuel.refresh_state()["status"] == "ok"
       and fuel.get_index("EE")["source"] == fuel.SOURCE_BULLETIN)
finally:
    fuel.urllib.request.urlopen = _real_urlopen

ok("🔴 the browser never calls the feed: the feed HOST appears in fuel.py only — the page names EuroOilWatch in its attribution and never as a URL",
   "eurooilwatch.com" not in open(os.path.join(ROOT, "frontend", "index.html"), encoding="utf-8").read().lower()
   and "eurooilwatch" not in open(os.path.join(ROOT, "map", "index.html"), encoding="utf-8").read().lower()
   and "eurooilwatch.com" in _src["fuel.py"]
   and "EU Weekly Oil Bulletin via EuroOilWatch" in open(os.path.join(ROOT, "frontend", "index.html"), encoding="utf-8").read())
ok("...and lookahead.py / derived.py never fetch — they read the stored row only",
   "fuel.refresh" not in open(os.path.join(BACKEND, "lookahead.py"), encoding="utf-8").read()
   and "fuel.refresh" not in _src["derived.py"] and "ensure_fresh" not in _src["derived.py"]
   and "fuel.get_index" in _src["derived.py"])

# =========================================================================== #
#  3. Settings — validation, the yard never touches the index, BAF nulls       #
# =========================================================================== #
reset_db()
ok("target validation: negative and non-numeric refused, unknown key refused",
   costing.validate_target({"rate_eur_per_t": -1}) and costing.validate_target({"rate_eur_per_km": "x"})
   and costing.validate_target({"rate_eur_per_hour": 1}) and not costing.validate_target({"rate_eur_per_t": 2.5}))
r = costing.set_target({"rate_eur_per_load": 45.0, "rate_eur_per_km": 1.5}, by="admin")
ok("set_target writes only the sent keys; €/t stays None", r["ok"] and r["target"]["rate_eur_per_load"] == 45.0
   and r["target"]["rate_eur_per_km"] == 1.5 and r["target"]["rate_eur_per_t"] is None)
r = costing.set_target({"rate_eur_per_load": None}, by="admin")
ok("...a sent null clears one rate and leaves the others", r["target"]["rate_eur_per_load"] is None and r["target"]["rate_eur_per_km"] == 1.5)
ok("...stored in the config table under key 'costing', stamped with who",
   db.query("SELECT updated_by FROM config WHERE tenant_id = ? AND key = ?", (db.current_tenant(), "costing"))[0]["updated_by"] == "admin")
ok("fuel validation: share over 100 refused, negative yard refused, a 3-letter country refused",
   costing.validate_fuel({"share_pct": 101}) and costing.validate_fuel({"yard_eur_per_l": -0.1})
   and costing.validate_fuel({"country": "EST"}) and not costing.validate_fuel({"share_pct": 30, "yard_eur_per_l": 1.4}))

_real_urlopen = fuel.urllib.request.urlopen
fuel.urllib.request.urlopen = _fake_urlopen
_calls["mode"] = "ok"
fuel.refresh("EE", sync=True)
fuel.urllib.request.urlopen = _real_urlopen
before = dict(fuel.get_index("EE"))
r = costing.set_fuel({"yard_eur_per_l": 1.41, "share_pct": 30}, by="planner")
ok("⭐ the yard PUT does NOT change the bulletin row", r["ok"] and dict(fuel.get_index("EE")) == before)
ok("...and the yard is stored separately from the index", costing.settings()["fuel"]["yard_eur_per_l"] == 1.41
   and costing.settings()["fuel"]["share_pct"] == 30.0)
ok("🔴 BAF is None with a share but NO base", costing.baf_pct(1.922) is None
   and costing.summary(fuel.get_index("EE"))["baf_reason"] == "no base")
costing.set_fuel({"share_pct": None})
ok("...None with neither", costing.baf_pct(1.922) is None and costing.summary(fuel.get_index("EE"))["baf_reason"] == "no share"
   or costing.summary(fuel.get_index("EE"))["baf_reason"] == "no base")
# lock the base by the same call confirm uses, then the formula
lk = costing.lock_baf_base_if_empty(fuel.get_index("EE"), by="planner")
ok("lock_baf_base_if_empty stamps the base from the index row: 1.922 of 2026-09-07",
   lk["locked"] is True and costing.settings()["fuel"]["baf_base_eur_per_l"] == 1.922
   and costing.settings()["fuel"]["baf_base_bulletin_date"] == "2026-09-07" and costing.settings()["fuel"]["baf_base_set_by"] == "planner")
ok("...with a base but NO share, BAF is still None (reason 'no share')",
   costing.baf_pct(1.922) is None and costing.summary(fuel.get_index("EE"))["baf_reason"] == "no share")
costing.set_fuel({"share_pct": 30})
ok("⭐ the formula: index 2.018 vs base 1.922 at 30 % share = (2.018/1.922 − 1) × 0.30 = +1.4984 %",
   abs(costing.baf_pct(2.018) - 0.014984) < 1e-5 and costing.baf_pct(1.922) == 0.0)
ok("...a falling index gives a NEGATIVE BAF — the surcharge is a rebate then",
   costing.baf_pct(1.80) < 0 and abs(costing.baf_pct(1.80) - ((1.80 / 1.922 - 1) * 0.3)) < 1e-6)
ok("adjust: quote × (1 + BAF); None when either side is missing",
   costing.adjust(1000.0, 0.014984) == 1014.98 and costing.adjust(None, 0.01) is None and costing.adjust(1000.0, None) is None)
ok("a second lock does NOT move the base (index unchanged in the row, but the call is a no-op regardless)",
   costing.lock_baf_base_if_empty({"eur_per_l": 2.5, "bulletin_date": "2026-10-01"})["locked"] is False
   and costing.settings()["fuel"]["baf_base_eur_per_l"] == 1.922)
costing.reset_baf_base(by="admin")
ok("reset_baf_base clears all four base fields", all(costing.settings()["fuel"][k] is None for k in
   ("baf_base_eur_per_l", "baf_base_bulletin_date", "baf_base_set_at", "baf_base_set_by")))
ok("...and locking with NO index row locks nothing — a base is never guessed",
   costing.lock_baf_base_if_empty(None)["locked"] is False and costing.lock_baf_base_if_empty({"eur_per_l": None})["locked"] is False
   and costing.settings()["fuel"]["baf_base_eur_per_l"] is None)

# resolve_rates: the all-or-nothing rule
tg = {"rate_eur_per_load": 50.0, "rate_eur_per_t": None, "rate_eur_per_km": 1.2}
ok("resolve_rates: no route rate ⇒ the target set, source 'target'",
   costing.resolve_rates({"per_load": None, "per_t": None, "per_km": None}, tg) == ({"per_load": 50.0, "per_t": None, "per_km": 1.2}, "target"))
ok("🔴 ...ONE typed route rate ⇒ the route's rates ALONE — the target's €/km is NOT mixed in",
   costing.resolve_rates({"per_load": None, "per_t": 2.0, "per_km": None}, tg) == ({"per_load": None, "per_t": 2.0, "per_km": None}, "route"))
ok("...neither ⇒ all None and source None (prints 'rate not set')",
   costing.resolve_rates({}, {k: None for k in costing.TARGET_FIELDS}) == ({"per_load": None, "per_t": None, "per_km": None}, None))

# =========================================================================== #
#  4. On the Look-ahead: target pricing, Quote + BAF, the account rows          #
# =========================================================================== #
reset_db()
access.set_current("planner123")
db.execute("INSERT INTO locations (id, name, lat, lon) VALUES (?, ?, ?, ?)", ("L1", "Pit", 58.5, 24.0))
db.execute("INSERT INTO locations (id, name, lat, lon) VALUES (?, ?, ?, ?)", ("L2", "Site", 58.6, 24.4))
for _rid in ("R1", "R2", "R3"):
    db.execute("INSERT INTO routes (id, origin_id, dest_id, ipt) VALUES (?, ?, ?, ?)", (_rid, "L1", "L2", "IPT 1"))
V8 = "Rigid 8-wheeler (32t)"      # payload 20 t


def _geom(rid, prof, leg, km, hr):
    db.execute("INSERT INTO route_geometry (tenant_id, route_id, vehicle_profile, leg, alt_index, "
               "geometry, distance_km, duration_hr) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
               ("default", rid, prof, leg, 0, "[[24,58.5],[24.4,58.6]]", km, hr))


for _rid in ("R1", "R2"):
    _geom(_rid, V8, "loaded", 30.0, 0.75)
    _geom(_rid, V8, "return", 30.0, 0.65)         # 60 km round trip, 5 cycles / 10 h
MONTHS = tuple(range(1, 25))


def _line(rid, sect, qty):
    main.save_matrix_row(main.MatrixRow(
        route_id=rid, discipline="earthworks", section_id=sect, material_type="Small aggregate",
        material_description=None, vehicle_type=V8, submitted_by="tester", unit="t",
        status="Pending", cells=[main.Cell(month_index=m, quantity=qty) for m in MONTHS], ipt="IPT1"))
    main.set_route_status(rid, main.StatusUpdate(status="Approved"), discipline="earthworks", section_id=sect)


_line("R1", "WS1", 4000.0)      # no route rate — will take the target
_line("R2", "WS2", 4000.0)      # a typed €/t only — must NOT mix with the target
_line("R3", "WS3", 900.0)       # unbaked, no rate — target's per-load/per-t only, partial
network.set_route_planning("R2", {"rate_eur_per_t"}, rate_eur_per_t=2.0)

res = main.list_forecast_days()
L = {l["section_id"]: l for l in res["lines"]}
ok("before any target: R1 and R3 print 'rate not set' (eur None), R2 prices from its own €/t",
   L["WS1"]["week_derived"]["eur"] is None and L["WS3"]["week_derived"]["eur"] is None
   and L["WS2"]["week_derived"]["eur"] is not None and L["WS2"]["context"]["rate_source"] == "route"
   and L["WS1"]["context"]["rate_source"] is None)
ok("...the response carries a costing block and totals.eur_target_lines = 0",
   isinstance(res.get("costing"), dict) and res["costing"]["target_set"] is False
   and res["totals"]["eur_target_lines"] == 0 and res["totals"]["eur_adj"] is None and res["totals"]["baf_pct"] is None)

costing.set_target({"rate_eur_per_load": 45.0, "rate_eur_per_km": 1.5}, by="admin")
res = main.list_forecast_days()
L = {l["section_id"]: l for l in res["lines"]}
wd1 = [d for d in L["WS1"]["days"] if datetime.date.fromisoformat(d["day_date"]).weekday() <= 4][0]["derived"]
exp = wd1["trips"] * 45.0 + wd1["trips"] * 60.0 * 1.5
ok("⭐ with a target: R1 prices at trips×45 + trips×60 km×1.5 (round trip), source 'target'",
   abs(wd1["eur"] - round(exp, 2)) < 0.011 and L["WS1"]["context"]["rate_source"] == "target"
   and L["WS1"]["context"]["rate_set"] is True and L["WS1"]["week_derived"]["rate_source"] == "target", str(wd1))
ok("...the route's TYPED rates are still exposed as route_rates (all None here) beside the resolved ones",
   L["WS1"]["context"]["route_rates"] == {"per_load": None, "per_t": None, "per_km": None}
   and L["WS1"]["context"]["rates"] == {"per_load": 45.0, "per_t": None, "per_km": 1.5})
wd2 = [d for d in L["WS2"]["days"] if datetime.date.fromisoformat(d["day_date"]).weekday() <= 4][0]["derived"]
ok("🔴 R2 with its own €/t is UNCHANGED by the target — no per-load or per-km term added",
   abs(wd2["eur"] - round(wd2["tonnes"] * 2.0, 2)) < 0.011 and L["WS2"]["context"]["rate_source"] == "route", str(wd2))
wd3 = [d for d in L["WS3"]["days"] if datetime.date.fromisoformat(d["day_date"]).weekday() <= 4][0]["derived"]
ok("R3 (unbaked) prices the target's per-load term only and says it is partial (per-km needs a distance)",
   abs(wd3["eur"] - round(wd3["trips"] * 45.0, 2)) < 0.011 and wd3["eur_partial"] is True
   and L["WS3"]["context"]["rate_source"] == "target")
T = res["totals"]
ok("totals: 3 priced lines, 2 of them at the target rate; still no BAF",
   T["eur_lines"] == 3 and T["eur_target_lines"] == 2 and T["eur_adj"] is None and T["baf_pct"] is None
   and all(d["derived"]["eur_adj"] is None for l in res["lines"] for d in l["days"]))
ok("...costing.target_set is now True and carries the typed target",
   res["costing"]["target_set"] is True and res["costing"]["target"]["rate_eur_per_load"] == 45.0)

# now the index, a share and a base — via the real Confirm path
_real_urlopen = fuel.urllib.request.urlopen
fuel.urllib.request.urlopen = _fake_urlopen
_calls["mode"] = "ok"
fuel.refresh("EE", sync=True)
fuel.urllib.request.urlopen = _real_urlopen
costing.set_fuel({"share_pct": 30}, by="planner")
res = main.list_forecast_days()
ok("index + share but no base yet: still no BAF, reason 'no base' — Quote only",
   res["totals"]["baf_pct"] is None and res["costing"]["baf_reason"] == "no base"
   and res["costing"]["index_eur_per_l"] == 1.922)

mi, wi = days.commit_bucket()
w1 = L["WS1"]["week"]
c1 = main.confirm_forecast_week(main.WeekConfirm(route_id="R1", month_index=mi, discipline="earthworks",
                                                   section_id="WS1", week_index=wi, confirmed_by="planner"))
ok("⭐ the FIRST Confirm week with an index row and no base LOCKS the base: 1.922 of 2026-09-07",
   c1.get("baf_base", {}).get("locked") is True and costing.settings()["fuel"]["baf_base_eur_per_l"] == 1.922
   and costing.settings()["fuel"]["baf_base_bulletin_date"] == "2026-09-07"
   and costing.settings()["fuel"]["baf_base_set_by"] == "planner", str(c1.get("baf_base")))
ok("...and the week is confirmed as before — the hook is a side effect, not a gate",
   c1["week"]["status"] == "confirmed")
# 🔴 caught while building: the hook first landed in set_week (its tail is identical to
# confirm_week's). Editing a week must never lock the base — only Confirm does.
costing.reset_baf_base(by="admin")
_w_other = [w for w in range(1, weeks.weeks_in_month(mi) + 1) if w != wi][0]
_e = weeks.set_week("R1", mi, "earthworks", "WS1", _w_other, planned_qty=1234.0)
ok("🔴 EDITING a week (set_week) does NOT lock the base — only Confirm does",
   not _e.get("error") and "baf_base" not in _e and costing.settings()["fuel"]["baf_base_eur_per_l"] is None, str(_e)[:80])
c1b = main.confirm_forecast_week(main.WeekConfirm(route_id="R1", month_index=mi, discipline="earthworks",
                                                    section_id="WS1", week_index=_w_other, confirmed_by="planner"))
ok("...and confirming that edited week locks it (1.922 again)",
   c1b.get("baf_base", {}).get("locked") is True and costing.settings()["fuel"]["baf_base_eur_per_l"] == 1.922)
# move the index, confirm ANOTHER line: the base must not follow
fuel.set_manual(2.018, "2026-09-14", by="admin")
c2 = main.confirm_forecast_week(main.WeekConfirm(route_id="R2", month_index=mi, discipline="earthworks",
                                                   section_id="WS2", week_index=wi, confirmed_by="planner"))
ok("🔴 a SECOND Confirm does not move the base, though the index has moved to 2.018",
   c2.get("baf_base", {}).get("locked") is False and costing.settings()["fuel"]["baf_base_eur_per_l"] == 1.922)

res = main.list_forecast_days()
L = {l["section_id"]: l for l in res["lines"]}
pct = (2.018 / 1.922 - 1) * 0.30
ok("⭐ now every priced day carries Quote AND Quote + BAF; BAF = (2.018/1.922 − 1) × 30 %",
   res["totals"].get("baf_pct") is not None and abs(res["totals"]["baf_pct"] - pct) < 1e-5
   and all(abs(d["derived"]["eur_adj"] - round(d["derived"]["eur"] * (1 + res["totals"]["baf_pct"]), 2)) < 0.011
           for l in res["lines"] for d in l["days"] if d["derived"]["eur"] is not None))
ok("🔴 the quote is never replaced: eur is the same figure as before the base existed",
   abs([d for d in L["WS1"]["days"] if datetime.date.fromisoformat(d["day_date"]).weekday() <= 4][0]["derived"]["eur"] - round(exp, 2)) < 0.011)
ok("...week and totals sum both figures, and totals.eur_adj > totals.eur when the index rose",
   L["WS1"]["week_derived"]["eur_adj"] is not None and res["totals"]["eur_adj"] is not None
   and res["totals"]["eur_adj"] > res["totals"]["eur"]
   and abs(res["totals"]["eur_adj"] - sum(l["week_derived"]["eur_adj"] for l in res["lines"])) < 0.05)
ok("...R2's own-rate line is surcharged too — BAF applies to any quote, typed or target",
   L["WS2"]["week_derived"]["eur_adj"] is not None and L["WS2"]["context"]["rate_source"] == "route")

# the page: account rows, the costing block, and the statement cost
pg = lookahead.page(bucket="commit")
# ---- the render fixture for render_frontend.js: a page WITH a target, an index, a base
# and a share, so the widget and the Quote + BAF figures render populated (timestamps scrubbed)
def _scrub(o):
    if isinstance(o, dict):
        return {k: ("<ts>" if k.endswith("_at") and isinstance(v, str) else _scrub(v)) for k, v in o.items()}
    if isinstance(o, list):
        return [_scrub(x) for x in o]
    return o
_fix_dir = os.path.join(HERE, "fixtures")
os.makedirs(_fix_dir, exist_ok=True)
with open(os.path.join(_fix_dir, "lookahead_page_fuel.json"), "w", encoding="utf-8") as _fh:
    _fh.write(json.dumps(_scrub(pg), indent=1, ensure_ascii=False))
ok("the fuel render fixture is written: 3 lines, a BAF, 2 target-priced lines, index 2.018",
   len(pg["commit"]["lines"]) == 3 and pg["costing"]["baf_pct"] is not None
   and pg["commit"]["totals"]["eur_target_lines"] == 2 and pg["costing"]["index_eur_per_l"] == 2.018)
ok("the page carries costing at top level and on commit", isinstance(pg.get("costing"), dict)
   and pg["commit"].get("costing", {}).get("baf_pct") == pg["costing"]["baf_pct"])
acct = pg["account"]
if acct.get("rows"):
    A = {r["section_id"]: r for r in acct["rows"]}
    ok("account rows: planned € from the target on R1 (source 'target'), from the route on R2, and a + BAF figure on both",
       A["WS1"]["rate_source"] == "target" and A["WS1"]["planned_eur"] is not None and A["WS1"]["planned_eur_adj"] is not None
       and A["WS2"]["rate_source"] == "route" and A["WS2"]["planned_eur_adj"] is not None
       and abs(A["WS1"]["planned_eur_adj"] - round(A["WS1"]["planned_eur"] * (1 + pct), 2)) < 0.011, str(A["WS1"]))
    ok("...and the account counts the lines priced at the target", acct["eur_target_lines"] == 2)
else:
    ok("(no account week in this bucket — the two account assertions did not run)", False, str(acct.get("week")))
    ok("(no account week in this bucket — the two account assertions did not run)", False)
import collections as _co
_orig_q, _orig_x = db.query, db.execute
_cnt = _co.Counter()
def _q(sql, *a, **k):
    _cnt["query"] += 1
    if "fuel_index" in sql or "FROM config" in sql:
        _cnt["cost"] += 1
    return _orig_q(sql, *a, **k)
def _x(sql, *a, **k):
    _cnt["execute"] += 1
    return _orig_x(sql, *a, **k)
db.query, db.execute = _q, _x
try:
    costing.invalidate(); _cnt.clear(); lookahead.page(bucket="commit"); c = dict(_cnt)
    ok(f"🔴 the page read costs the costing slice at most a handful of extra statements ({c.get('cost', 0)} config/fuel reads) and writes nothing",
       c.get("execute", 0) == 0 and c.get("cost", 0) <= 8, str(c))
finally:
    db.query, db.execute = _orig_q, _orig_x

# =========================================================================== #
#  5. Endpoints — bodies, roles, the admin token                               #
# =========================================================================== #
pl = main.get_costing()
ok("GET /api/costing: target, fuel, index and summary", set(pl) >= {"target", "fuel", "index", "summary"}
   and pl["index"]["eur_per_l"] == 2.018 and pl["summary"]["baf_pct"] is not None)
os.environ["ADMIN_TOKEN"] = "s3cret"
try:
    e = _raises(lambda: main.put_costing_target(main.TargetRatesIn(rate_eur_per_t=3.0, updated_by="x"), token="wrong"))
    ok("PUT /api/admin/costing/target refuses a bad admin token", getattr(e, "status_code", None) == 403)
    e = _raises(lambda: main.put_fuel_manual(main.ManualIndexIn(eur_per_l=1.9, bulletin_date="2026-09-14"), token="wrong"))
    ok("PUT /api/admin/fuel-index/manual refuses a bad admin token", getattr(e, "status_code", None) == 403)
    e = _raises(lambda: main.reset_fuel_base(token="wrong"))
    ok("POST /api/admin/fuel-index/reset-base refuses a bad admin token", getattr(e, "status_code", None) == 403)
    e = _raises(lambda: main.refresh_fuel_index(token="wrong"))
    ok("POST /api/admin/fuel-index/refresh refuses a bad admin token", getattr(e, "status_code", None) == 403)
    r = main.put_costing_target(main.TargetRatesIn(rate_eur_per_t=3.0, updated_by="x"), token="s3cret")
    ok("...and writes with the right one — only the field SENT (€/t), the others untouched",
       r["target"]["rate_eur_per_t"] == 3.0 and r["target"]["rate_eur_per_load"] == 45.0 and r["target"]["rate_eur_per_km"] == 1.5)
    e = _raises(lambda: main.put_costing_target(main.TargetRatesIn(rate_eur_per_t=-3.0), token="s3cret"))
    ok("...a negative rate is a 400", getattr(e, "status_code", None) == 400)
finally:
    os.environ.pop("ADMIN_TOKEN", None)
access.set_current("submitter123")
e = _raises(lambda: main.put_fuel_settings(main.FuelSettingsIn(share_pct=20)))
ok("PUT /api/fuel-index/settings: a submitter is refused (403) — planner or admin only",
   getattr(e, "status_code", None) == 403)
access.set_current("planner123")
r = main.put_fuel_settings(main.FuelSettingsIn(yard_eur_per_l=1.39, updated_by="p"))
ok("...a planner writes the yard price; the share it did not send is untouched",
   r["fuel"]["yard_eur_per_l"] == 1.39 and r["fuel"]["share_pct"] == 30.0)
e = _raises(lambda: main.put_fuel_settings(main.FuelSettingsIn(share_pct=150)))
ok("...share 150 is a 400", getattr(e, "status_code", None) == 400)
_real_urlopen = fuel.urllib.request.urlopen
fuel.urllib.request.urlopen = _fake_urlopen
try:
    _calls["mode"] = "500"; _calls["n"] = 0
    db.execute("UPDATE fuel_index SET last_attempt_at = ? WHERE country = ?", ("2026-09-01T00:00:00Z", "EE"))
    r = main.get_fuel_index(lazy=1, sync=1)
    ok("🔴 GET /api/fuel-index with the feed down returns 200-shaped data: the last good row + refresh.status 'unavailable', not a 500",
       r["index"]["eur_per_l"] == 2.018 and r["index"]["refresh"]["status"] == "unavailable" and _calls["n"] == 1)
    _calls["n"] = 0
    main.get_fuel_index(lazy=1, sync=1)
    ok("...a second GET within 12 h does NOT hit the feed again (a failing feed is not hammered)", _calls["n"] == 0)
    _calls["n"] = 0
    main.get_fuel_index(lazy=0)
    ok("...lazy=0 never touches the feed", _calls["n"] == 0)
    _calls["mode"] = "ok"
    r = main.refresh_fuel_index(sync=1)
    ok("POST /api/admin/fuel-index/refresh (token unset ⇒ open, as every admin endpoint is today — C11) fetches and stores 1.922",
       r["index"]["eur_per_l"] == 1.922 and r["index"]["source"] == fuel.SOURCE_BULLETIN)
    r = main.put_fuel_manual(main.ManualIndexIn(eur_per_l=1.95, bulletin_date="2026-09-14", updated_by="a"))
    ok("PUT manual writes source 'manual' and the widget state says so", r["index"]["source"] == "manual" and r["index"]["eur_per_l"] == 1.95)
    r = main.reset_fuel_base(updated_by="a")
    ok("POST reset-base clears the base; the next confirm may lock again", r["fuel"]["baf_base_eur_per_l"] is None
       and r["summary"]["baf_reason"] == "no base")
finally:
    fuel.urllib.request.urlopen = _real_urlopen

# =========================================================================== #
#  6. Export — the columns, the Fuel sheet, and what is NOT on the supplier's  #
#     sheet                                                                    #
# =========================================================================== #
# a base again so the export has a BAF to print
costing.lock_baf_base_if_empty(fuel.get_index("EE"), by="planner")
pg = lookahead.page(bucket="commit")
xb = export.build_xlsx(pg)
import openpyxl as _ox
wb = _ox.load_workbook(io.BytesIO(xb))
ws = wb["Commit week"]
hdr = [ws.cell(row=1, column=j).value for j in range(1, ws.max_column + 1)]
ok("XLSX: '€ source' and '€ + BAF' columns follow '€'",
   "€" in hdr and hdr.index("€ source") == hdr.index("€") + 1 and hdr.index("€ + BAF") == hdr.index("€") + 2, str(hdr))
col_src, col_adj, col_eur = hdr.index("€ source") + 1, hdr.index("€ + BAF") + 1, hdr.index("€") + 1
vals = [(ws.cell(row=i, column=col_eur).value, ws.cell(row=i, column=col_src).value, ws.cell(row=i, column=col_adj).value)
        for i in range(2, ws.max_row + 1)]
ok("...every priced row says 'target' or 'route', and € + BAF = € × (1 + BAF) on each",
   all(s in ("target", "route") for e_, s, a in vals if e_ is not None)
   and any(s == "target" for _, s, _ in vals) and any(s == "route" for _, s, _ in vals)
   and all(abs(a - round(e_ * (1 + pg["costing"]["baf_pct"]), 2)) < 0.011 for e_, s, a in vals if e_ is not None), str(vals[:3]))
wf = wb["Fuel"]
ftxt = " ".join(str(wf.cell(row=i, column=j).value or "") for i in range(1, wf.max_row + 1) for j in (1, 2))
ok("the Fuel sheet names the index, bulletin date, base, share, BAF % and the attribution",
   "Diesel index" in ftxt and "2026-09-14" in ftxt and "BAF base" in ftxt and "Fuel share" in ftxt
   and "EuroOilWatch" in ftxt and "BAF %" in ftxt)
ok("🔴 ...and NOT the planner's own numbers: no yard price, no target rate on the supplier's workbook",
   "Yard" not in ftxt and "Target €" not in ftxt and "1.39" not in ftxt)
pb = export.build_pdf(pg)
if os.environ.get("RBE_PDF_OUT"):            # look at it (lesson 24): RBE_PDF_OUT=/path python3 …
    with open(os.environ["RBE_PDF_OUT"], "wb") as _pf:
        _pf.write(pb)
_txt = subprocess.run(["pdftotext", "-layout", "-", "-"], input=pb, capture_output=True).stdout.decode("utf-8") if shutil.which("pdftotext") else pb.decode("latin-1")
ok("the PDF names the diesel index, its bulletin and the BAF formula in the footer",
   "Diesel EE" in _txt and "1.950/L" in _txt and "BAF" in _txt and "fuel share" in _txt)
ok("...and marks a target-priced line and prints +BAF under the quote", "target" in _txt and "+BAF" in _txt)
ok("...the footer no longer says € is only from the route", "€ not printed where no contract rate is typed on the route." not in " ".join(export.FOOTER))
# no quote ⇒ diesel columns present, quote_adj blank
costing.set_target({k: None for k in costing.TARGET_FIELDS}, by="admin")
network.set_route_planning("R2", {"rate_eur_per_t"}, rate_eur_per_t=None)
pg0 = lookahead.page(bucket="commit")
xb0 = export.build_xlsx(pg0)
wb0 = _ox.load_workbook(io.BytesIO(xb0))
ws0 = wb0["Commit week"]
ok("🔴 no quote anywhere ⇒ the export still has the € + BAF column and the Fuel sheet, and every € + BAF is blank — no fake €",
   "€ + BAF" in [ws0.cell(row=1, column=j).value for j in range(1, ws0.max_column + 1)] and "Fuel" in wb0.sheetnames
   and all(ws0.cell(row=i, column=col_adj).value is None for i in range(2, ws0.max_row + 1))
   and pg0["commit"]["totals"]["eur"] is None and pg0["commit"]["totals"]["eur_adj"] is None)
pb0 = export.build_pdf(pg0)
_txt0 = subprocess.run(["pdftotext", "-layout", "-", "-"], input=pb0, capture_output=True).stdout.decode("utf-8") if shutil.which("pdftotext") else pb0.decode("latin-1")
ok("...and the PDF still names the diesel index with no € column", "Diesel EE" in _txt0 and "€ WEEK" not in _txt0)

# =========================================================================== #
#  7. What must NOT exist                                                      #
# =========================================================================== #
_map = open(os.path.join(ROOT, "map", "index.html"), encoding="utf-8").read().lower()
ok("🔴 no fuel on the public map: no 'fuel-index', no 'baf', no '€/l'; 'diesel' only where it always was (the CO₂e caption, once)",
   "fuel-index" not in _map and "fuel index" not in _map and re.search(r"\bbaf\b", _map) is None
   and "€/l" not in _map and _map.count("diesel") == 1 and "round-trip diesel" in _map)
ok("no new upload endpoint: main.py has no UploadFile / File( import", "UploadFile" not in _src["main.py"] and "from fastapi import File" not in _src["main.py"])
ok("no paid or station API named in fuel.py", all(x not in _src["fuel.py"].lower().split("what this does not do")[1]
   for x in ("oilpriceapi.com", "fuelo.", "alexela.", "circlek.")) and "energy.ec.europa.eu" not in _src["fuel.py"])
ok("no L/100 km, no €/h standing rate, no UNDER_COST as a NAME anywhere in the backend (docstrings may say so)",
   all(not any(x in n for n in _tok[f][0] for x in ("l_per_100", "per_100km", "under_cost", "rate_eur_per_hour", "eur_per_hour"))
       for f in _tok))
ok("the settings PUT never carries the base — locking is Confirm's job, resetting is the admin's",
   "baf_base" not in " ".join(costing.FUEL_FIELDS) and "baf_base_eur_per_l" not in _src["main.py"].split("class FuelSettingsIn")[1].split("class ManualIndexIn")[0])
ok("the widget's mounts exist in the page source: FuelWidget defined once, mounted three times (Account, Commit, Config)",
   len(re.findall(r"function FuelWidget\(", open(os.path.join(ROOT, "frontend", "index.html"), encoding="utf-8").read())) == 1
   and len(re.findall(r"<FuelWidget\b", open(os.path.join(ROOT, "frontend", "index.html"), encoding="utf-8").read())) == 3)

print()
for f in FAIL:
    print("  FAIL:", f)
print(f"\n{PASS} passed, {len(FAIL)} failed")
sys.exit(1 if FAIL else 0)
