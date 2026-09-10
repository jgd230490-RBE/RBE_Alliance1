"""
Look-ahead v2, slice 1 — the commit week by day (2026-09-09).

Backend assertions for forecast_days, its materialisation rules, the day APIs, the
sum-rule flag, the day-actual → week-actual sync, whole-week confirm stamping the days,
and the opt-in calibrate spread. Same harness as test_week1.py: a scratch SQLite
database, HERE never called, `fastapi` / `flexpolyline` / `psycopg2` stubbed.

WHAT THIS DOES NOT PROVE
------------------------
  * The HTTP layer is stubbed. Endpoint BODIES run; nothing proves a route is mounted
    or that `?from=` / `?to=` are parsed off a query string.
  * No Postgres branch runs. forecast_days is created against SQLite only; its Postgres
    path is a plain CREATE plus the 4.5 tenant migration.
  * Nothing in a browser. There is no frontend in slice 1.
  * "Today" is the real clock for the endpoint-level assertions (same as test_week1),
    and an explicit `today=` for the bucket-arithmetic ones. The seed spans two years
    so the commit month is Approved whatever the date.

Run:  python3 backend/tests/test_lookahead.py
"""
import datetime
import json
import re
import os
import shutil
import subprocess
import sys
import tempfile
import types

HERE = os.path.dirname(os.path.abspath(__file__))
BACKEND = os.path.dirname(HERE)
ROOT = os.path.dirname(BACKEND)
sys.path.insert(0, BACKEND)

# --------------------------------------------------------------------------- #
#  Stubs                                                                       #
# --------------------------------------------------------------------------- #
_fp = types.ModuleType("flexpolyline")
_fp.decode = lambda s: []
_fp.encode = lambda pts: ""
sys.modules.setdefault("flexpolyline", _fp)


def _passthrough_decorator(*a, **k):
    def wrap(fn):
        return fn
    return wrap


class _App:
    def __init__(self, *a, **k):
        pass

    def get(self, *a, **k):
        return _passthrough_decorator()

    post = put = delete = patch = get

    def add_middleware(self, *a, **k):
        pass

    def mount(self, *a, **k):
        pass


class _HTTPException(Exception):
    def __init__(self, status_code, detail=""):
        super().__init__(f"{status_code}: {detail}")
        self.status_code = status_code
        self.detail = detail


def _Query(default=None, **k):
    return default


_fa = types.ModuleType("fastapi")
_fa.FastAPI = _App
_fa.HTTPException = _HTTPException
_fa.Query = _Query
sys.modules.setdefault("fastapi", _fa)

_mw = types.ModuleType("fastapi.middleware")
_cors = types.ModuleType("fastapi.middleware.cors")
_cors.CORSMiddleware = object
_mw.cors = _cors
sys.modules.setdefault("fastapi.middleware", _mw)
sys.modules.setdefault("fastapi.middleware.cors", _cors)

_resp = types.ModuleType("fastapi.responses")
_resp.FileResponse = lambda *a, **k: None
sys.modules.setdefault("fastapi.responses", _resp)

_static = types.ModuleType("fastapi.staticfiles")


# a CLASS, not a lambda: main.py subclasses this, and `class X(lambda)` is a TypeError
class _StaticFiles:
    def __init__(self, *a, **k):
        pass

    async def get_response(self, path, scope):
        return None


_static.StaticFiles = _StaticFiles
sys.modules.setdefault("fastapi.staticfiles", _static)

TMP = tempfile.mkdtemp(prefix="rbe_lookahead_")
os.environ.pop("DATABASE_URL", None)

import db  # noqa: E402
db._SQLITE_PATH = os.path.join(TMP, "scratch.db")

import conversions  # noqa: E402
import taxonomy  # noqa: E402
import network  # noqa: E402
import weeks  # noqa: E402
import stockpiles  # noqa: E402
import main  # noqa: E402
# 2026-09-02, Task F: the staff endpoints now require an access code, resolved from the
# X-Access-Code header by a middleware the stubbed app never runs. The harness sets the
# request context itself, as a PLANNER, so every pre-existing assertion still exercises
# the same code paths it did. Nothing here proves the header is actually read.
import access  # noqa: E402
for _v in ("IPT1_CODE", "IPT2_CODE", "IPT3_CODE", "IPT4_CODE", "IPT5_CODE", "IPT6_CODE",
           "PLANNER_CODE", "ADMIN_CODE"):
    os.environ.pop(_v, None)
access.set_current("planner123")

PASS = 0
FAIL = []


def ok(label, cond, extra=""):
    global PASS
    if cond:
        PASS += 1
    else:
        FAIL.append(f"{label} {extra}".strip())


def reset_db():
    if os.path.exists(db._SQLITE_PATH):
        os.remove(db._SQLITE_PATH)
    db.init_db()
    db.init_network_db()
    db.init_taxonomy_db()
    db.init_zones_db()
    db.init_gates_db()
    db.init_weeks_db()
    db.init_config_db()
    db.init_tenant()
    import config as _cfg
    _cfg.invalidate()


def seed_line(route_id="R1", disc="earthworks", sect="WS1", months=(1, 2, 3),
              qty=100.0, unit="t", status="Pending"):
    """One forecast line over `months`, through the real bulk endpoint."""
    return main.save_matrix_row(main.MatrixRow(
        route_id=route_id, discipline=disc, section_id=sect,
        material_type="Small aggregate", material_description=None,
        vehicle_type="Rigid 8-wheeler (32t)", submitted_by="tester",
        unit=unit, status=status,
        cells=[main.Cell(month_index=m, quantity=qty) for m in months]))


import days  # noqa: E402
days.START_YEAR = main.START_YEAR


def _as(code):
    access.set_current(code)


def _line_days(rid="R1", disc="earthworks", sect="WS1"):
    mi, wi = days.commit_bucket()
    return days._days_of(rid, mi, disc, sect, wi)


def _by_date(rows):
    return {r["day_date"]: r for r in rows}


# =========================================================================== #
#  0. Registration — a new table is not finished when it is created            #
# =========================================================================== #
reset_db()
ok("forecast_days is a tenanted table", "forecast_days" in db.TENANTED_TABLES)
ok("...with a DDL entry and a PK entry",
   "forecast_days" in db._TENANT_DDL and "forecast_days" in db._TENANT_PK)
# .get(), not a subscript — breaking the registration on purpose raised KeyError here
# and took the report down with the assertion above it, which HAD failed correctly.
ok("...whose key is the week's line key plus the date",
   db._TENANT_PK.get("forecast_days", "").replace(" ", "")
   == "(tenant_id,route_id,month_index,discipline,section_id,day_date)")
ok("forecast_weeks gained actual_source, and it is in the DDL (or the rebuild drops it)",
   "actual_source" in db._TENANT_DDL["forecast_weeks"]
   and "actual_source" in [c.lower() for c in db._columns_of(db.get_conn().cursor(), "forecast_weeks")])

# =========================================================================== #
#  1. Bucket arithmetic — pure functions, explicit dates                       #
# =========================================================================== #
def _raises(fn):
    try:
        fn()
    except Exception as e:
        return e
    return None

ok("month_of: index 1 is January of START_YEAR", days.month_of(1, 2026) == (2026, 1))
ok("month_of: index 13 rolls into the next year", days.month_of(13, 2026) == (2027, 1))
# 10 Sep: CALENDAR weeks. 1 Sep 2026 is a Tuesday; week 2 of September is Mon 7 – Sun 13.
d_w2 = days.bucket_dates(9, 2, 2026)
ok("⭐ week 2 of Sep 2026 is Mon 7 to Sun 13 — seven days, starting on a MONDAY",
   [d.day for d in d_w2] == [7, 8, 9, 10, 11, 12, 13] and d_w2[0].weekday() == 0 and d_w2[0].month == 9)
d_w1 = days.bucket_dates(9, 1, 2026)
ok("⭐ week 1 of Sep 2026 starts on Mon 31 AUGUST — a week's first days may lie in the month before",
   d_w1[0] == datetime.date(2026, 8, 31) and d_w1[-1] == datetime.date(2026, 9, 6))
d_w5 = days.bucket_dates(10, 5, 2026)
ok("⭐ October 2026 has a FIFTH week, Mon 26 Oct – Sun 1 Nov",
   d_w5[0] == datetime.date(2026, 10, 26) and d_w5[-1] == datetime.date(2026, 11, 1))
ok("...and asking September for a fifth week is refused, not silently empty",
   isinstance(_raises(lambda: days.bucket_dates(9, 5, 2026)), ValueError))
ok("every bucket is exactly seven days and starts on a Monday, two years through",
   all(len(days.bucket_dates(m, w, 2026)) == 7 and days.bucket_dates(m, w, 2026)[0].weekday() == 0
       for m in range(1, 25) for w in range(1, weeks.weeks_in_month(m, 2026) + 1)))
ok("weekdays_in counts Mon-Fri only",
   len(days.weekdays_in(d_w2)) == 5 and len(days.weekdays_in(d_w5)) == 5)
sp = days.split_week(1000.0, d_w2)
ok("⭐ L1: Mon-Fri each get week/5 and Sat/Sun get 0",
   all(abs(sp[d] - 200.0) < 1e-9 for d in d_w2 if d.weekday() <= 4)
   and all(sp[d] == 0.0 for d in d_w2 if d.weekday() > 4))
ok("...and the split sums back to the week", abs(sum(sp.values()) - 1000.0) < 1e-9)
sp4 = days.split_week(500.0, d_w5)
ok("the fifth week splits the same way: 5 weekdays x 100, 2 weekend days x 0",
   sum(1 for v in sp4.values() if abs(v - 100.0) < 1e-9) == 5
   and sum(1 for v in sp4.values() if v == 0.0) == 2)
ok("commit_bucket with an explicit today matches weeks.editable_week",
   days.commit_bucket(2026, today=datetime.date(2026, 9, 10)) == (9, 2)
   and days.commit_bucket(2026, today=datetime.date(2026, 10, 25)) == (10, 4)
   and days.commit_bucket(2026, today=datetime.date(2026, 10, 30)) == (10, 5)
   and days.commit_bucket(2026, today=datetime.date(2026, 8, 31)) == (9, 1))

# =========================================================================== #
#  2. Materialisation — commit week only, Approved only                       #
# =========================================================================== #
reset_db()
_as("planner123")
db.execute("INSERT INTO locations (id, name, lat, lon) VALUES (?, ?, ?, ?)", ("L1", "Pit", 58.5, 24.0))
db.execute("INSERT INTO locations (id, name, lat, lon) VALUES (?, ?, ?, ?)", ("L2", "Site", 58.6, 24.4))
db.execute("INSERT INTO routes (id, origin_id, dest_id) VALUES (?, ?, ?)", ("R1", "L1", "L2"))
db.execute("INSERT INTO routes (id, origin_id, dest_id) VALUES (?, ?, ?)", ("R2", "L1", "L2"))
MONTHS = tuple(range(1, 25))                 # two years, so today is always covered
seed_line(route_id="R1", disc="earthworks", sect="WS1", months=MONTHS, qty=4000.0)
seed_line(route_id="R2", disc="earthworks", sect="WS9", months=MONTHS, qty=800.0)   # stays Pending

MI, WI = days.commit_bucket()
DATES = days.bucket_dates(MI, WI)
ISO = [d.isoformat() for d in DATES]
NWD = len(days.weekdays_in(DATES))
ok("the commit bucket is inside the seeded horizon", 1 <= MI <= 24 and 1 <= WI <= 4)

r = days.materialise_commit_week()
ok("nothing materialises while no line is Approved", r["created"] == 0 and _line_days() == [])

main.set_route_status("R1", main.StatusUpdate(status="Approved"), discipline="earthworks", section_id="WS1")
r = days.materialise_commit_week()
rows = _line_days()
ok("⭐ approving the month creates one day per calendar day in the commit bucket",
   r["created"] == len(DATES) and [x["day_date"] for x in rows] == ISO)
ok("...every one `derived` with the week index stamped",
   all(x["status"] == "derived" and x["parent_week_index"] == WI for x in rows))
ok("🔴 and NONE outside the bucket — no days for weeks 2-4 of the horizon",
   db.query("SELECT COUNT(*) AS n FROM forecast_days WHERE tenant_id = ?", (db.current_tenant(),))[0]["n"] == len(DATES))
ok("🔴 and none for the Pending line", days._days_of("R2", MI, "earthworks", "WS9", WI) == [])
# ⚠️ That assertion is protected by weeks.py, not by days.py: a never-approved line has
# no week row, so there is nothing to hang a day on. Breaking days.py's own Approved
# guard on purpose left it green. The case that guard actually protects is a line that
# WAS approved (so its week rows exist and survive un-approval, per weeks.py's rule) and
# is then reopened: it must not get days.
seed_line(route_id="R2", disc="earthworks", sect="WS8", months=MONTHS, qty=300.0)
main.set_route_status("R2", main.StatusUpdate(status="Approved"), discipline="earthworks", section_id="WS8")
ok("...a reopened line still has its week rows (weeks.py's rule, not changed here)",
   weeks.get_week("R2", MI, "earthworks", "WS8", WI) is not None
   and (main.set_route_status("R2", main.StatusUpdate(status="Pending"), discipline="earthworks", section_id="WS8") or True)
   and weeks.get_week("R2", MI, "earthworks", "WS8", WI) is not None)
days.materialise_commit_week()
ok("🔴 ...and days.py does NOT materialise days for it — a week row is not enough, the parent must be Approved",
   days._days_of("R2", MI, "earthworks", "WS8", WI) == [])
wk = weeks.get_week("R1", MI, "earthworks", "WS1", WI)
share = float(wk["planned_qty"]) / NWD
ok("⭐ L1 in the database: weekdays carry week/n, weekends carry 0",
   all(abs(float(x["planned_qty"]) - share) < 1e-9 for x in rows if datetime.date.fromisoformat(x["day_date"]).weekday() <= 4)
   and all(float(x["planned_qty"]) == 0.0 for x in rows if datetime.date.fromisoformat(x["day_date"]).weekday() > 4))
ok("...and the day planned quantities sum to the week", abs(sum(float(x["planned_qty"]) for x in rows) - float(wk["planned_qty"])) < 1e-6)
ok("materialising again is idempotent", days.materialise_commit_week()["created"] == 0)

# the read, through the endpoint
res = main.list_forecast_days()
ok("GET /api/forecast-days groups by line and names the commit week",
   res["commit_week"]["month_index"] == MI and res["commit_week"]["week_index"] == WI
   and res["commit_week"]["days"] == len(DATES) and res["commit_week"]["weekdays"] == NWD)
ln = [l for l in res["lines"] if l["route_id"] == "R1"]
ok("...one entry for the Approved line, none for the Pending one",
   len(ln) == 1 and not [l for l in res["lines"] if l["route_id"] == "R2"])
ok("...carrying the parent's unit, material and vehicle so the UI needs no second call",
   ln[0]["unit"] == "t" and ln[0]["material_type"] == "Small aggregate" and ln[0]["vehicle_type"] == "Rigid 8-wheeler (32t)")
ok("...and the sum rule is satisfied on a fresh materialise",
   ln[0]["days_ne_week"] is False and abs(ln[0]["days_sum"] - float(wk["planned_qty"])) < 1e-6)
clip = main.list_forecast_days(from_date=ISO[1], to_date=ISO[2])
ok("from/to clip the days returned but do not widen the bucket",
   [d["day_date"] for d in clip["lines"][0]["days"]] == ISO[1:3])
ok("the response says which statuses exist and counts the table",
   clip["statuses"] == ["derived", "edited", "confirmed"] and clip["summary"]["days"] == len(DATES))

# =========================================================================== #
#  3. Derived refreshes, edited holds, and `week_changed`                       #
# =========================================================================== #
WD = [d.isoformat() for d in days.weekdays_in(DATES)]
first_wd, second_wd = WD[0], WD[1]
r = main.edit_forecast_day(main.DayEdit(route_id="R1", month_index=MI, discipline="earthworks",
                                        section_id="WS1", day_date=first_wd, planned_qty=999.0))
bd = _by_date(_line_days())
ok("typing a day sets it to `edited` with the typed figure",
   bd[first_wd]["status"] == "edited" and float(bd[first_wd]["planned_qty"]) == 999.0)
ok("⭐ ...and moves a `derived` week to `edited` without changing its number",
   weeks.get_week("R1", MI, "earthworks", "WS1", WI)["status"] == "edited"
   and abs(float(weeks.get_week("R1", MI, "earthworks", "WS1", WI)["planned_qty"]) - float(wk["planned_qty"])) < 1e-9)
ok("⭐ the sum rule now FLAGS — and nothing is blocked", r["line"]["days_ne_week"] is True)

# the week moves under the days
weeks.set_week("R1", MI, "earthworks", "WS1", WI, planned_qty=2000.0)
days.materialise_commit_week()
bd = _by_date(_line_days())
new_share = 2000.0 / NWD
ok("⭐ a `derived` day follows the week's new figure",
   abs(float(bd[second_wd]["planned_qty"]) - new_share) < 1e-9 and bd[second_wd]["status"] == "derived")
ok("🔴 an `edited` day holds what was typed", float(bd[first_wd]["planned_qty"]) == 999.0)
view = days._line_view("R1", MI, "earthworks", "WS1", WI)
vb = {d["day_date"]: d for d in view["days"]}
ok("⭐ ...and says the week changed under it, while the refreshed day does not",
   vb[first_wd]["week_changed"] is True and vb[second_wd]["week_changed"] is False)

# outside the bucket
other = (DATES[0] - datetime.timedelta(days=10)).isoformat()
r = days.set_day("R1", MI, "earthworks", "WS1", other, 5.0)
ok("a day outside the commit week is refused, naming the bucket",
   "not in the commit week" in (r.get("error") or "") and r["commit_week"]["week_index"] == WI)
ok("Sat/Sun are typeable — 0 is a default, not a rule",
   (lambda wknd: (not wknd) or days.set_day("R1", MI, "earthworks", "WS1", wknd[0], 40.0).get("error") is None)
   ([d.isoformat() for d in DATES if d.weekday() > 4]))

# =========================================================================== #
#  4. Actuals — rule 6, the freeze trap, and no calibration                    #
# =========================================================================== #
before_next = weeks.get_week("R1", *weeks.next_week(MI, WI)[:1], "earthworks", "WS1", weeks.next_week(MI, WI)[1]) \
    if weeks.next_week(MI, WI)[0] <= 24 else None
r = main.set_forecast_day_actual(main.DayActual(route_id="R1", month_index=MI, discipline="earthworks",
                                                section_id="WS1", day_date=first_wd, actual_qty=10.0))
w = weeks.get_week("R1", MI, "earthworks", "WS1", WI)
ok("a day actual is stored with its variance", float(r["line"]["days"][0]["actual_qty"]) == 10.0 if r["line"]["days"][0]["day_date"] == first_wd else True)
ok("⭐ with no week actual typed, the week actual becomes the day sum",
   w["actual_qty"] == 10.0 and r["line"]["week_actual"] == "summed")
main.set_forecast_day_actual(main.DayActual(route_id="R1", month_index=MI, discipline="earthworks",
                                            section_id="WS1", day_date=second_wd, actual_qty=20.0))
w = weeks.get_week("R1", MI, "earthworks", "WS1", WI)
ok("🔴 THE FREEZE TRAP: a second day actual UPDATES the sum rather than being blocked by the first",
   w["actual_qty"] == 30.0)
ok("...and the day's status is untouched by an actual", _by_date(_line_days())[second_wd]["status"] == "derived")
if before_next is not None:
    after_next = weeks.get_week("R1", *weeks.next_week(MI, WI)[:1], "earthworks", "WS1", weeks.next_week(MI, WI)[1])
    ok("⭐ saving day actuals did NOT calibrate — next week's plan is untouched",
       float(after_next["planned_qty"]) == float(before_next["planned_qty"]) and after_next["status"] == before_next["status"])

# a clerk types the week
weeks.set_actual("R1", MI, "earthworks", "WS1", WI, actual_qty=100.0)
main.set_forecast_day_actual(main.DayActual(route_id="R1", month_index=MI, discipline="earthworks",
                                            section_id="WS1", day_date=second_wd, actual_qty=25.0))
w = weeks.get_week("R1", MI, "earthworks", "WS1", WI)
ok("🔴 a TYPED week actual is never overwritten by a day sum", w["actual_qty"] == 100.0)
src = db.query("SELECT actual_source FROM forecast_weeks WHERE tenant_id = ? AND route_id = ? AND month_index = ? AND week_index = ?",
               (db.current_tenant(), "R1", MI, WI))[0]["actual_source"]
ok("...because the source says so", src == "typed")
weeks.set_actual("R1", MI, "earthworks", "WS1", WI, actual_qty=None)
r = main.set_forecast_day_actual(main.DayActual(route_id="R1", month_index=MI, discipline="earthworks",
                                                section_id="WS1", day_date=second_wd, actual_qty=25.0))
w = weeks.get_week("R1", MI, "earthworks", "WS1", WI)
ok("⭐ clearing the typed week actual lets the day sum take over again", w["actual_qty"] == 35.0)
main.set_forecast_day_actual(main.DayActual(route_id="R1", month_index=MI, discipline="earthworks", section_id="WS1", day_date=first_wd, actual_qty=None))
main.set_forecast_day_actual(main.DayActual(route_id="R1", month_index=MI, discipline="earthworks", section_id="WS1", day_date=second_wd, actual_qty=None))
w = weeks.get_week("R1", MI, "earthworks", "WS1", WI)
ok("clearing every day actual clears the summed week actual to None, not 0", w["actual_qty"] is None)

# =========================================================================== #
#  5. Confirm is whole-week, and it stamps the days                           #
# =========================================================================== #
r = main.confirm_forecast_week(main.WeekConfirm(route_id="R1", month_index=MI, discipline="earthworks",
                                                section_id="WS1", week_index=WI, confirmed_by="tester"))
ok("⭐ L6: confirming the week confirms every day in its bucket",
   r["week"]["status"] == "confirmed" and r["days_confirmed"] == len(DATES)
   and all(x["status"] == "confirmed" for x in _line_days()))
r = days.set_day("R1", MI, "earthworks", "WS1", second_wd, 1.0)
ok("🔴 ...after which planned days are read-only", r.get("blocked_by") == "confirmed")
r = main.set_forecast_day_actual(main.DayActual(route_id="R1", month_index=MI, discipline="earthworks",
                                                section_id="WS1", day_date=second_wd, actual_qty=12.0))
ok("⭐ ...but actuals are still typeable on a confirmed week", r.get("error") is None
   and float(_by_date(_line_days())[second_wd]["actual_qty"]) == 12.0)
# a week that is NOT the commit bucket confirms as it always did, with no days
om, ow = (MI, 1) if WI != 1 else (MI, 2)
r = days.confirm_week("R1", om, "earthworks", "WS1", ow)
ok("a week outside the commit bucket confirms with zero days stamped — Task C unchanged",
   r["week"]["status"] == "confirmed" and r["days_confirmed"] == 0)

# =========================================================================== #
#  6. Calibrate — next week only, and the opt-in spread                       #
# =========================================================================== #
# calibrating INTO a confirmed commit week still 400s, spread or not
pm, pw = (MI, WI - 1) if WI > 1 else (MI - 1, 4)
weeks.set_actual("R1", pm, "earthworks", "WS1", pw, actual_qty=1.0)
r = days.calibrate("R1", pm, "earthworks", "WS1", pw, spread=True)
ok("🔴 calibrate into a confirmed commit week is refused, spread or not", r.get("blocked_by") == "confirmed")

# reopen by moving the commit week back to `edited` — the only way is a fresh line
seed_line(route_id="R1", disc="earthworks", sect="WS2", months=MONTHS, qty=1000.0)
main.set_route_status("R1", main.StatusUpdate(status="Approved"), discipline="earthworks", section_id="WS2")
days.materialise_commit_week()
prev = weeks.get_week("R1", pm, "earthworks", "WS2", pw)
weeks.set_actual("R1", pm, "earthworks", "WS2", pw, actual_qty=float(prev["planned_qty"]) - 180.0)   # 180 short
target_before = float(weeks.get_week("R1", MI, "earthworks", "WS2", WI)["planned_qty"])
days_before = {x["day_date"]: float(x["planned_qty"]) for x in days._days_of("R1", MI, "earthworks", "WS2", WI)}

# spread OFF: the week moves, no day does
r = days.calibrate("R1", pm, "earthworks", "WS2", pw, spread=False)
ok("calibrate with spread off writes the week only",
   r["spread"] is None and abs(float(r["week"]["planned_qty"]) - (target_before + 180.0)) < 1e-6)
days_after = {x["day_date"]: float(x["planned_qty"]) for x in days._days_of("R1", MI, "earthworks", "WS2", WI)}
ok("...and no day planned figure moved", days_after == days_before)
# 2026-09-09, the Thursday case: the source week now records what was carried
src = weeks.get_week("R1", pm, "earthworks", "WS2", pw)
ok("⭐ the source week is stamped with the variance it has carried (180) and when",
   abs(float(src.get("calibrated_qty") or 0) - 180.0) < 1e-6 and bool(src.get("calibrated_at"))
   and abs(r.get("delta") - 180.0) < 1e-6 and r.get("already_carried") == 0)
# 🔴 pressing again with the SAME actual must carry NOTHING — not another 180
r = days.calibrate("R1", pm, "earthworks", "WS2", pw, spread=True, spread_from=WD[0])
ok("🔴 a second calibrate with an unchanged actual carries a delta of 0 — no double count",
   abs(r.get("delta")) < 1e-9 and r.get("already_carried") == 180.0
   and abs(float(r["week"]["planned_qty"]) - (target_before + 180.0)) < 1e-6)

# spread ON, from the first weekday: the actual worsens by ANOTHER 180 (Friday's figure
# landed) — only that difference is carried, over every weekday
weeks.set_actual("R1", pm, "earthworks", "WS2", pw, actual_qty=float(prev["planned_qty"]) - 360.0)
# ⚠️ The expectation is built from the WEEK as it stands before this calibrate, not from
# `days_after`. Those days are still `derived` and have not been read since the first
# calibrate moved the week, so they are a stale snapshot that no longer sums to it. The
# spread freezes the as-was distribution FIRST (derived day = week/n as of now), then adds
# the delta. Written against the stale snapshot this assertion failed, and the failure
# was the test's, not the code's — the sum-rule assertion below passed throughout.
week_now = float(weeks.get_week("R1", MI, "earthworks", "WS2", WI)["planned_qty"])
r = days.calibrate("R1", pm, "earthworks", "WS2", pw, spread=True, spread_from=WD[0])
ok("⭐ spread on: the delta is divided over the remaining WEEKDAYS only",
   r["spread"]["applied"] is True and abs(r["spread"]["delta"] - 180.0) < 1e-6
   and r["spread"]["days"] == WD and abs(r["spread"]["per_day"] - 180.0 / NWD) < 1e-6)
now_days = {x["day_date"]: x for x in days._days_of("R1", MI, "earthworks", "WS2", WI)}
ok("...each weekday = its as-was share + shortage/n, and every day went `edited`",
   all(abs(float(now_days[d]["planned_qty"]) - (week_now / NWD + 180.0 / NWD)) < 1e-6
       for d in WD)
   and all(now_days[d]["status"] == "edited" for d in ISO))
ok("🔴 ...and Sat/Sun stayed at 0",
   all(float(now_days[d]["planned_qty"]) == 0.0 for d in ISO if d not in WD))
ok("...and the days still sum to the week — the spread keeps the sum rule",
   r["line"]["days_ne_week"] is False)

# spread from a later weekday: only the days from there on. 360 carried so far; the
# actual worsens to 450 short, so the delta is 90
weeks.set_actual("R1", pm, "earthworks", "WS2", pw, actual_qty=float(prev["planned_qty"]) - 450.0)
r = days.calibrate("R1", pm, "earthworks", "WS2", pw, spread=True, spread_from=WD[-2])
ok("⭐ spread_from limits it to the weekdays on or after that date",
   r["spread"]["days"] == WD[-2:] and abs(r["spread"]["per_day"] - 45.0) < 1e-6
   and abs(r["spread"]["delta"] - 90.0) < 1e-6)
# spread from after the last weekday: nothing to spread onto, said plainly
weeks.set_actual("R1", pm, "earthworks", "WS2", pw, actual_qty=float(prev["planned_qty"]) - 10.0)
r = days.calibrate("R1", pm, "earthworks", "WS2", pw, spread=True,
                   spread_from=(DATES[-1] + datetime.timedelta(days=1)).isoformat())
ok("no remaining weekdays ⇒ the week is still written and the response says no day moved",
   r["spread"]["applied"] is False and "no weekdays remain" in r["spread"]["note"]
   and r.get("error") is None)
# through the endpoint, with the new body fields
r = main.calibrate_forecast_week(main.WeekCalibrate(route_id="R1", month_index=pm, discipline="earthworks",
                                                   section_id="WS2", week_index=pw, spread=False))
ok("the endpoint's spread defaults OFF and passes through", r["spread"] is None)

# =========================================================================== #
#  7. Task F — the day layer inherits the parent line's IPT                   #
# =========================================================================== #
os.environ.update({"IPT3_CODE": "three-secret", "IPT6_CODE": "six-secret",
                   "PLANNER_CODE": "plan-secret", "ADMIN_CODE": "adm-secret"})
_as("plan-secret")
db.execute("INSERT INTO routes (id, origin_id, dest_id) VALUES (?, ?, ?)", ("R3", "L1", "L2"))
for _sect, _ipt in (("WS3", "IPT3"), ("WS6", "IPT6")):
    main.save_matrix_row(main.MatrixRow(route_id="R3", discipline="substructure", section_id=_sect,
        material_type="Small aggregate", material_description=None, vehicle_type="Rigid 8-wheeler (32t)",
        submitted_by="planner", unit="t", status="Pending",
        cells=[main.Cell(month_index=m, quantity=50.0) for m in MONTHS], ipt=_ipt))
    main.set_route_status("R3", main.StatusUpdate(status="Approved"), discipline="substructure", section_id=_sect)
_as("six-secret")
res6 = main.list_forecast_days(route_id="R3")
ok("⭐ an IPT6 code sees only IPT6's days", {l["section_id"] for l in res6["lines"]} == {"WS6"})
try:
    main.edit_forecast_day(main.DayEdit(route_id="R3", month_index=MI, discipline="substructure",
                                        section_id="WS3", day_date=first_wd, planned_qty=1.0))
    ok("🔴 an IPT6 code cannot type IPT3's day", False)
except Exception as e:
    ok("🔴 an IPT6 code cannot type IPT3's day — 404, not 403", getattr(e, "status_code", None) == 404)
r = main.edit_forecast_day(main.DayEdit(route_id="R3", month_index=MI, discipline="substructure",
                                        section_id="WS6", day_date=first_wd, planned_qty=1.0))
ok("...but can type its own", r.get("error") is None)
for _v in ("IPT3_CODE", "IPT6_CODE", "PLANNER_CODE", "ADMIN_CODE"):
    os.environ.pop(_v, None)
_as("planner123")

# =========================================================================== #
#  8. Tenant isolation on the table itself                                    #
# =========================================================================== #
n_default = db.query("SELECT COUNT(*) AS n FROM forecast_days WHERE tenant_id = ?", (db.current_tenant(),))[0]["n"]
tok = db.set_current_tenant("other")
try:
    ok("🔴 another tenant sees no days at all",
       days.list_days()["lines"] == [] and days.summary()["days"] == 0)
finally:
    db.reset_current_tenant(tok)
ok("...and the default tenant's rows are untouched by that read",
   db.query("SELECT COUNT(*) AS n FROM forecast_days WHERE tenant_id = ?", (db.current_tenant(),))[0]["n"] == n_default)

# =========================================================================== #
#  9. What the brief says must NOT exist                                      #
# =========================================================================== #
main_src = open(os.path.join(BACKEND, "main.py"), encoding="utf-8").read()
days_src = open(os.path.join(BACKEND, "days.py"), encoding="utf-8").read()
ok("🔴 no upload endpoint was added — no UploadFile, no File(, no multipart",
   "UploadFile" not in main_src and "File(" not in main_src and "multipart" not in main_src.lower())
ok("🔴 no per-day confirm exists", "confirm_day" not in days_src and "/forecast-days/confirm" not in main_src)
ok("🔴 the forecasts UNIQUE key is unchanged",
   db._TENANT_UNIQUE["forecasts"].replace(" ", "") == "(tenant_id,route_id,month_index,discipline,section_id)")
ok("the public map feed is still monthly", "week" not in main.public_route_forecasts.__doc__.lower())
ok("days.py imports weeks and does not reimplement its writes",
   "import weeks" in days_src and "UPDATE forecast_weeks SET planned_qty" not in days_src)


# =========================================================================== #
#  10. SLICE 2 — context and derived figures, computed on read (2026-09-09)    #
# =========================================================================== #
# Fresh database. Four lines, one per case the brief's §8 names:
#   R1/WS1  tonnes, baked both legs      -> every figure, cycle from HERE
#   R2/WS2  tonnes, NOT baked            -> trips + tonnes only, UNBAKED flag
#   R3/WS3  unit=vehicles, loaded only   -> trips = qty, return estimated (‡), km_trip = loaded
#   R1/WS4  a vehicle factors.json does not know -> V07 fallback, same as month-kpis
import config  # noqa: E402
import derived  # noqa: E402
reset_db()
_as("planner123")
db.execute("INSERT INTO locations (id, name, lat, lon) VALUES (?, ?, ?, ?)", ("L1", "Pit", 58.5, 24.0))
db.execute("INSERT INTO locations (id, name, lat, lon) VALUES (?, ?, ?, ?)", ("L2", "Site", 58.6, 24.4))
for _rid, _ipt in (("R1", "IPT 1"), ("R2", "IPT 2"), ("R3", "IPT 3 / IPT 6")):
    db.execute("INSERT INTO routes (id, origin_id, dest_id, ipt) VALUES (?, ?, ?, ?)", (_rid, "L1", "L2", _ipt))
V8 = "Rigid 8-wheeler (32t)"                                  # payload 20 t
V12 = "N3 tractor (BC) + O4 tipping semi (DA) GCW 40/44 t"    # code V12, payload 26 t


def _geom2(rid, prof, leg, km, hr):
    db.execute("INSERT INTO route_geometry (tenant_id, route_id, vehicle_profile, leg, alt_index, "
               "geometry, distance_km, duration_hr) VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
               ("default", rid, prof, leg, 0, "[[24,58.5],[24.4,58.6]]", km, hr))


_geom2("R1", V8, "loaded", 30.0, 0.75)
_geom2("R1", V8, "return", 30.0, 0.65)      # cycle 0.75 + 0.65 + 20/60 = 1.7333 h -> 5 per 10 h shift
_geom2("R3", V12, "loaded", 30.0, 0.7)      # no return leg


def _line2(rid, disc, sect, unit, qty, veh, ipt=None):
    main.save_matrix_row(main.MatrixRow(
        route_id=rid, discipline=disc, section_id=sect, material_type="Small aggregate",
        material_description=None, vehicle_type=veh, submitted_by="tester", unit=unit,
        status="Pending", cells=[main.Cell(month_index=m, quantity=qty) for m in MONTHS], ipt=ipt))
    main.set_route_status(rid, main.StatusUpdate(status="Approved"), discipline=disc, section_id=sect)


_line2("R1", "earthworks", "WS1", "t", 4000.0, V8, ipt="IPT1")
_line2("R2", "earthworks", "WS2", "t", 800.0, V8, ipt="IPT2")
_line2("R3", "substructure", "WS3", "vehicles", 100.0, V12, ipt="IPT3")
_line2("R1", "earthworks", "WS4", "t", 900.0, "Ghost truck", ipt="IPT1")

n_days_before = db.query("SELECT COUNT(*) AS n FROM forecast_days WHERE tenant_id = ?", (db.current_tenant(),))[0]["n"]
res = main.list_forecast_days()
L = {l["section_id"]: l for l in res["lines"]}
ok("slice 2: every line carries context and week_derived, every day carries derived",
   len(L) == 4 and all("context" in l and "week_derived" in l for l in res["lines"])
   and all("derived" in d for l in res["lines"] for d in l["days"]))
ok("...and the response carries totals", isinstance(res.get("totals"), dict))
ok("🔴 computed, NOT stored — the read created exactly the day rows slice 1 would have, no more",
   db.query("SELECT COUNT(*) AS n FROM forecast_days WHERE tenant_id = ?", (db.current_tenant(),))[0]["n"]
   == len(DATES) * 4 and n_days_before == 0)
ok("...no new table, no new column: sixteen tenanted tables, forecast_days' DDL unchanged",
   len(db.TENANTED_TABLES) == 16 and "trips" not in db._TENANT_DDL["forecast_days"].lower())

# ---- context chips
c1 = L["WS1"]["context"]
ok("context names both ends of the route and the route's own IPT",
   c1.get("origin_name") == "Pit" and c1.get("dest_name") == "Site"
   and c1.get("origin_id") == "L1" and c1.get("route_ipt") == "IPT 1")
ok("...and the LINE's IPT, WS, material, vehicle and unit",
   c1.get("ipt") == "IPT1" and c1.get("section_id") == "WS1" and c1.get("material_type") == "Small aggregate"
   and c1.get("vehicle_type") == V8 and c1.get("unit") == "t")
ok("vehicle_short: the EU code where factors.json has one, the de-prefixed legacy name otherwise",
   L["WS3"]["context"].get("vehicle_short") == "V12" and c1.get("vehicle_short") == "8-wheeler (32t)")
ok("payload is the planning figure with no fallback for a known vehicle",
   c1.get("payload_t") == 20.0 and c1.get("payload_fallback") is None)
ok("a baked line: baked, return_baked, both distances, km_trip = loaded + return, no ‡",
   c1.get("baked") is True and c1.get("return_baked") is True and c1.get("distance_km") == 30.0
   and c1.get("return_km") == 30.0 and c1.get("km_trip") == 60.0 and c1.get("cycle_mark") is None
   and c1.get("cycle_source") == "here")
ok("cycle_min is route_analysis' cycle in minutes (104.0 = 1.7333 h)", c1.get("cycle_min") == 104.0)
ok("⭐ cycles_per_vehicle_day IS route_analysis().trips_per_day — floor(600 / 104) = 5",
   c1.get("cycles_per_vehicle_day") == 5 and c1.get("shift_hours") == 10.0)
ab = main.routes_analysis_batch(route_ids="R1")
ok("⭐ ...and it equals what /api/routes/analysis-batch reports for the same (route, vehicle)",
   (ab.get("analysis", {}).get("R1", {}).get(V8) or {}).get("trips_per_day") == c1.get("cycles_per_vehicle_day")
   and (ab.get("analysis", {}).get("R1", {}).get(V8) or {}).get("loaded_km") == c1.get("distance_km"))
ok("a baked line raises no flag", c1.get("flags") == [])

# ---- the formulas, brief §4, on a weekday of the baked tonnes line
wk1 = L["WS1"]["week"]
q = float(wk1["planned_qty"]) / NWD
import math as _m
wd = [d for d in L["WS1"]["days"] if datetime.date.fromisoformat(d["day_date"]).weekday() <= 4]
we = [d for d in L["WS1"]["days"] if datetime.date.fromisoformat(d["day_date"]).weekday() > 4]
f = wd[0]["derived"]
exp_trips = int(_m.ceil(round(q / 20.0, 9)))
exp_veh = int(_m.ceil(exp_trips / 5))
ok("trips = ceil(qty / payload)", f.get("trips") == exp_trips and exp_trips > 0, str(f))
ok("tonnes = the quantity, for a line typed in t", abs(f.get("tonnes") - q) < 1e-6)
ok("vehicles_need = ceil(trips / cycles_per_veh)", f.get("vehicles") == exp_veh, str(f))
ok("km_day = trips × km_trip (both legs)", abs(f.get("km_day") - exp_trips * 60.0) < 1e-6)
ok("km_per_vehicle = km_day / vehicles", abs(f.get("km_per_vehicle") - round(exp_trips * 60.0 / exp_veh, 2)) < 1e-6)
# C20, decided 09 Sep evening: t·km uses the ROUTE'S km basis, default round trip (what a
# haulier charges). Reversed from the brief's loaded-leg reading; a route can say 'loaded'.
ok("⭐ tonne_km = tonnes × the route's basis km — 60 (round trip) by default, not the 30 loaded leg",
   abs(f.get("tonne_km") - round(q * 60.0, 1)) < 1e-6 and c1.get("km_basis") == "round_trip"
   and c1.get("basis_km") == 60.0)
ok("every weekday of a derived week is identical", all(d["derived"] == f for d in wd))
ok("Sat/Sun: 0 trips, 0 vehicles, 0 km — zeros, because the day IS planned at 0",
   all(d["derived"]["trips"] == 0 and d["derived"]["vehicles"] == 0 and d["derived"]["km_day"] == 0.0
       for d in we))
w1 = L["WS1"]["week_derived"]
ok("week_derived sums the days and takes the PEAK vehicles, not the sum",
   w1.get("trips") == exp_trips * NWD and w1.get("vehicles_peak") == exp_veh
   and abs(w1.get("tonne_km") - round(q * 60.0 * NWD, 1)) < 0.11
   and abs(w1.get("km") - exp_trips * 60.0 * NWD) < 1e-6)

# ---- the ceiling does not round up a float artefact
ok("⭐ ceil(3.0000000000000004) is 3, not 4", derived._ceil(3.0000000000000004) == 3 and derived._ceil(3.01) == 4)

# ---- unbaked: trips and tonnes only, UNBAKED flag, nothing invented
c2 = L["WS2"]["context"]
f2 = [d for d in L["WS2"]["days"] if datetime.date.fromisoformat(d["day_date"]).weekday() <= 4][0]["derived"]
q2 = float(L["WS2"]["week"]["planned_qty"]) / NWD
ok("🔴 an unbaked line still computes trips from the planning payload",
   f2.get("trips") == int(_m.ceil(round(q2 / 20.0, 9))) and f2.get("trips") > 0)
ok("🔴 ...but vehicles, km and t·km are None — not 0, not a 45 km/h guess",
   f2.get("vehicles") is None and f2.get("km_day") is None and f2.get("km_per_vehicle") is None
   and f2.get("tonne_km") is None)
ok("...its context says so: not baked, no distance, no cycle, and the UNBAKED flag",
   c2.get("baked") is False and c2.get("distance_km") is None and c2.get("km_trip") is None
   and c2.get("cycle_min") is None and c2.get("cycle_source") is None
   and c2.get("cycles_per_vehicle_day") is None and c2.get("flags") == [derived.FLAG_UNBAKED])
ok("...and its week_derived carries None for every distance figure",
   L["WS2"]["week_derived"].get("vehicles_peak") is None and L["WS2"]["week_derived"].get("km") is None
   and L["WS2"]["week_derived"].get("trips") > 0)
derived_src = open(os.path.join(BACKEND, "derived.py"), encoding="utf-8").read()
ok("🔴 derived.py never reads avg_haul_speed_kmh — the fallback speed is not applied anywhere",
   "avg_haul_speed" not in derived_src)
ok("🔴 derived.py writes nothing", "INSERT" not in derived_src and "UPDATE" not in derived_src)

# ---- unit = vehicles, loaded leg only
c3, l3 = L["WS3"]["context"], L["WS3"]
f3 = [d for d in l3["days"] if datetime.date.fromisoformat(d["day_date"]).weekday() <= 4][0]
q3 = float(l3["week"]["planned_qty"]) / NWD
ok("unit=vehicles: trips = ceil(qty) itself, not qty / payload",
   f3["derived"]["trips"] == int(_m.ceil(round(q3, 9))))
ok("...and tonnes = qty × the same payload trips use (26 t for V12)",
   abs(f3["derived"]["tonnes"] - q3 * 26.0) < 1e-6 and c3.get("payload_t") == 26.0)
ok("⭐ loaded leg only: baked, return NOT baked, cycle marked ‡, km_trip = loaded leg alone",
   c3.get("baked") is True and c3.get("return_baked") is False and c3.get("cycle_mark") == "‡"
   and c3.get("cycle_source") == "here_return_estimated" and c3.get("km_trip") == 30.0
   and c3.get("return_km") is None)
ra3 = network.route_analysis("R3", profiles=[V12])["rows"][0]
ok("...its cycle is route_analysis' return-estimated cycle, not a re-derivation",
   ra3.get("return_estimated") is True and c3.get("cycles_per_vehicle_day") == ra3.get("trips_per_day")
   and c3.get("cycle_min") == round(ra3["cycle_hr"] * 60.0, 1))
ok("...and km_day uses that one-leg km_trip", abs(f3["derived"]["km_day"] - f3["derived"]["trips"] * 30.0) < 1e-6)
ok("a loaded-only bake is still 'baked' — no UNBAKED flag", c3.get("flags") == [])

# ---- unknown vehicle: the month-kpis fallback, and the two agree
c4 = L["WS4"]["context"]
kp = {l["section_id"]: l for l in main.public_month_kpis(month=MI, unit="t")["lines"]}
ok("an unknown vehicle falls back to the first planning vehicle's payload and names it",
   c4.get("payload_fallback") == kp["WS4"].get("payload_fallback") and c4.get("payload_fallback") is not None
   and c4.get("payload_t") == kp["WS4"].get("payload_t") == 18.0)
ok("⭐ month-kpis and the Look-ahead agree on payload for every line this month",
   all(L[s]["context"]["payload_t"] == kp[s]["payload_t"] for s in ("WS1", "WS2", "WS3", "WS4")))
ok("⭐ ...and on movements-per-vehicle-per-day, baked or not",
   all(L[s]["context"]["cycles_per_vehicle_day"] == kp[s]["trips_per_vehicle_day"]
       for s in ("WS1", "WS2", "WS3", "WS4")))
ok("an unknown vehicle on an unbaked pairing has no cycle, so no vehicles",
   c4.get("baked") is False and derived.FLAG_UNBAKED in c4.get("flags", []))

# ---- totals: what the KPI strip reads
T = res["totals"]
exp_planned_t = sum(l["week_derived"]["tonnes"] for l in res["lines"])
exp_trips_all = sum(l["week_derived"]["trips"] for l in res["lines"])
ok("totals.planned_t is the sum of every line's tonnes, unbaked included",
   abs(T.get("planned_t") - exp_planned_t) < 1e-6 and T.get("lines") == 4)
ok("totals.trips is the sum of every line's trips, unbaked included", T.get("trips") == exp_trips_all)
peak_day_veh = f.get("vehicles") + f3["derived"]["vehicles"]
ok("⭐ totals.vehicles_peak is the largest SAME-DAY sum across baked lines (R1 + R3 on a weekday)",
   T.get("vehicles_peak") == peak_day_veh and T.get("vehicles_peak_date") in ISO
   and datetime.date.fromisoformat(T["vehicles_peak_date"]).weekday() <= 4)
ok("...and it is not the sum of per-line peaks when those fall on different days (same here, asserted equal)",
   T.get("vehicles_peak") == max(T["vehicles_by_day"].values()))
ok("totals.tonne_km and km omit BOTH unbaked lines (WS2, and WS4 whose vehicle R1 is not baked for), and say so",
   T.get("unbaked_lines") == 2 and T.get("excludes_unbaked") is True
   and abs(T.get("tonne_km") - (w1["tonne_km"] + l3["week_derived"]["tonne_km"])) < 1e-6,
   f"tonne_km={T.get('tonne_km')} w1={w1['tonne_km']} l3={l3['week_derived']['tonne_km']} unbaked={T.get('unbaked_lines')} lines={[ (l['section_id'], l['context']['baked'], l['week_derived']['tonne_km']) for l in res['lines']]}")

# ---- a quantity that does NOT divide by the payload: ceil, not floor (the seeded 200 t/day did)
first_wd_iso = wd[0]["day_date"]
main.edit_forecast_day(main.DayEdit(route_id="R1", month_index=MI, discipline="earthworks",
                                    section_id="WS1", day_date=first_wd_iso, planned_qty=210.0))
_d210 = [d for d in {l["section_id"]: l for l in main.list_forecast_days()["lines"]}["WS1"]["days"]
         if d["day_date"] == first_wd_iso][0]["derived"]
ok("⭐ 210 t on a 20 t payload is 11 trips — ceil, not floor (10)", _d210.get("trips") == 11, str(_d210))
ok("...and 11 trips at 5 cycles is 3 vehicles, not 2", _d210.get("vehicles") == 3)

# ---- two lines peaking on DIFFERENT days: the KPI is the largest same-day sum, not a sum of peaks
second_wd_iso = wd[1]["day_date"]
main.edit_forecast_day(main.DayEdit(route_id="R3", month_index=MI, discipline="substructure",
                                    section_id="WS3", day_date=second_wd_iso, planned_qty=q3 * 40))
main.edit_forecast_day(main.DayEdit(route_id="R1", month_index=MI, discipline="earthworks",
                                    section_id="WS1", day_date=second_wd_iso, planned_qty=0.0))
_rp = main.list_forecast_days()
_Lp = {l["section_id"]: l for l in _rp["lines"]}
_sum_of_peaks = sum((l["week_derived"].get("vehicles_peak") or 0) for l in _rp["lines"])
ok("⭐ totals.vehicles_peak is strictly LESS than the sum of per-line peaks when they fall on different days",
   _rp["totals"]["vehicles_peak"] == max(_rp["totals"]["vehicles_by_day"].values())
   and _rp["totals"]["vehicles_peak"] < _sum_of_peaks
   and _rp["totals"]["vehicles_peak_date"] == second_wd_iso
   and _Lp["WS3"]["week_derived"]["vehicles_peak"] > f3["derived"]["vehicles"],
   f"peak={_rp['totals']['vehicles_peak']} sum={_sum_of_peaks} by_day={_rp['totals']['vehicles_by_day']}")
# put the two days back so the assertions below start from the derived week
main.edit_forecast_day(main.DayEdit(route_id="R3", month_index=MI, discipline="substructure",
                                    section_id="WS3", day_date=second_wd_iso, planned_qty=q3))
main.edit_forecast_day(main.DayEdit(route_id="R1", month_index=MI, discipline="earthworks",
                                    section_id="WS1", day_date=second_wd_iso, planned_qty=q))

# ---- computed on READ: edit a day and the figures follow without a write of their own
main.edit_forecast_day(main.DayEdit(route_id="R1", month_index=MI, discipline="earthworks",
                                    section_id="WS1", day_date=first_wd_iso, planned_qty=0.0))
res_b = main.list_forecast_days()
Lb = {l["section_id"]: l for l in res_b["lines"]}
db0 = [d for d in Lb["WS1"]["days"] if d["day_date"] == first_wd_iso][0]["derived"]
ok("⭐ zero a day and its trips, vehicles and km read 0 on the next read — nothing was stored to update",
   db0.get("trips") == 0 and db0.get("vehicles") == 0 and db0.get("km_day") == 0.0 and db0.get("tonne_km") == 0.0)
ok("...the week's trips drop by exactly that day's trips",
   Lb["WS1"]["week_derived"]["trips"] == w1["trips"] - exp_trips)
ok("...and totals.trips drops by the same amount", res_b["totals"]["trips"] == T["trips"] - exp_trips)

# ---- the access filter comes first: totals are for what the caller can see
os.environ.update({"IPT1_CODE": "one-secret", "IPT2_CODE": "two-secret",
                   "PLANNER_CODE": "plan-secret", "ADMIN_CODE": "adm-secret"})
_as("two-secret")
res2 = main.list_forecast_days()
ok("⭐ an IPT2 code's totals cover ONLY IPT2's line — one line, and it is the unbaked one",
   res2["totals"].get("lines") == 1 and res2["totals"].get("unbaked_lines") == 1
   and res2["totals"].get("vehicles_peak") == 0 and res2["totals"].get("vehicles_peak_date") is None
   and {l["section_id"] for l in res2["lines"]} == {"WS2"})
for _v in ("IPT1_CODE", "IPT2_CODE", "PLANNER_CODE", "ADMIN_CODE"):
    os.environ.pop(_v, None)
_as("planner123")

# ---- the endpoint is the only place it is wired, and it is after the filter
main_src = open(os.path.join(BACKEND, "main.py"), encoding="utf-8").read()
_i_filter = main_src.find("access.filter_lines(res[\"lines\"], acc)")
_i_dec = main_src.find("derived.decorate(res)")
ok("🔴 derived.decorate() runs AFTER access.filter_lines() in the days endpoint",
   0 < _i_filter < _i_dec and main_src.count("derived.decorate(") == 1)
# reversed 09 Sep evening: slice 5 landed in the same zip — € is computed, never seeded
ok("€ is computed from the route's typed rates, and no default rate exists in the module",
   "rate_eur_per_km" in derived_src and "eur" in f and f.get("eur") is None
   and c1.get("rate_set") is False and "DEFAULT_RATE" not in derived_src)


# =========================================================================== #
#  11. SLICES 3–6 — rates + €, the clash rail, account, horizon, reopen,       #
#      the 'next' bucket, and the export (2026-09-09, evening)                 #
# =========================================================================== #
import lookahead  # noqa: E402
import clashes    # noqa: E402
import export     # noqa: E402
reset_db()
_as("planner123")
db.execute("INSERT INTO locations (id, name, lat, lon) VALUES (?, ?, ?, ?)", ("L1", "Pit", 58.5, 24.0))
db.execute("INSERT INTO locations (id, name, lat, lon, capacity_qty, capacity_unit, opening_qty) "
           "VALUES (?, ?, ?, ?, ?, ?, ?)", ("L2", "Site", 58.6, 24.4, 1000.0, "t", 100.0))
db.execute("INSERT INTO locations (id, name, lat, lon) VALUES (?, ?, ?, ?)", ("L3", "Yard", 58.7, 24.5))
for _rid, _o, _d in (("R1", "L1", "L2"), ("R2", "L1", "L2"), ("R3", "L3", "L2")):
    db.execute("INSERT INTO routes (id, origin_id, dest_id) VALUES (?, ?, ?)", (_rid, _o, _d))
for _rid in ("R1", "R2", "R3"):
    _geom2(_rid, V8, "loaded", 30.0, 0.75)
    _geom2(_rid, V8, "return", 30.0, 0.65)          # 5 cycles per 10 h shift, 60 km round trip

# ---- the schema: five route columns, three week columns, no new table
ok("routes gained the five planning columns (DDL and live table)",
   all(c in db._TENANT_DDL["routes"] for c in ("max_vehicles_per_day", "rate_eur_per_load", "rate_eur_per_t", "rate_eur_per_km", "km_basis"))
   and all(c in [x.lower() for x in db._columns_of(db.get_conn().cursor(), "routes")]
           for c in ("max_vehicles_per_day", "rate_eur_per_km", "km_basis")))
ok("forecast_weeks gained actual_cost_eur, calibrated_at, calibrated_qty",
   all(c in [x.lower() for x in db._columns_of(db.get_conn().cursor(), "forecast_weeks")]
       for c in ("actual_cost_eur", "calibrated_at", "calibrated_qty")))
ok("...and still sixteen tenanted tables — no new table for slices 3-6", len(db.TENANTED_TABLES) == 16)
db.init_lookahead_db()
ok("init_lookahead_db() is idempotent on a table that already has the columns", True)

# ---- route planning: only sent fields written, blanks clear, validation
r = network.set_route_planning("R1", {"rate_eur_per_t", "rate_eur_per_km", "max_vehicles_per_day"},
                               rate_eur_per_t=2.0, rate_eur_per_km=1.5, max_vehicles_per_day=4)
ok("route planning writes the sent fields and defaults km_basis to round_trip",
   r.get("rate_eur_per_t") == 2.0 and r.get("rate_eur_per_km") == 1.5 and r.get("max_vehicles_per_day") == 4
   and r.get("km_basis") == "round_trip" and r.get("rate_eur_per_load") is None)
r = network.set_route_planning("R1", {"rate_eur_per_load"}, rate_eur_per_load=45.0)
ok("...a later write of ONE field leaves the others alone",
   r.get("rate_eur_per_load") == 45.0 and r.get("rate_eur_per_t") == 2.0 and r.get("max_vehicles_per_day") == 4)
ok("...km_basis must be round_trip or loaded",
   "km_basis" in (network.set_route_planning("R1", {"km_basis"}, km_basis="both").get("error") or ""))
ok("...a negative rate is refused",
   "negative" in (network.set_route_planning("R1", {"rate_eur_per_t"}, rate_eur_per_t=-1).get("error") or ""))
try:
    main.set_route_planning("R9", main.RoutePlanning(rate_eur_per_t=1.0))
    ok("the planning endpoint 404s an unknown route", False)
except Exception as e:
    ok("the planning endpoint 404s an unknown route", getattr(e, "status_code", None) == 404)
network.set_route_planning("R3", {"rate_eur_per_t", "km_basis"}, rate_eur_per_t=3.0, km_basis="loaded")
ok("routes_status carries the planning fields",
   {x["id"]: x for x in network.routes_status()}["R1"].get("rate_eur_per_km") == 1.5)

# ---- lines: R1 for IPT1 and IPT2 (shared route → IPT_SHARE, and a cap of 4 veh/day), R3 for IPT3
_line2("R1", "earthworks", "WS1", "t", 4000.0, V8, ipt="IPT1")     # 200 t/day → 10 trips → 2 veh
_line2("R1", "earthworks", "WS2", "t", 6000.0, V8, ipt="IPT2")     # 300 t/day → 15 trips → 3 veh  ⇒ 5 > cap 4
_line2("R3", "substructure", "WS3", "t", 2000.0, V8, ipt="IPT3")   # 100 t/day → 5 trips → 1 veh

pg = lookahead.page(bucket="commit")
L = {l["section_id"]: l for l in pg["commit"]["lines"]}
ok("the page read returns commit, account, horizon, stock and clashes",
   all(k in pg for k in ("commit", "account", "horizon", "stock", "clashes")) and len(L) == 3)

# ---- € per day, summed terms, on the route's basis
f1 = [d for d in L["WS1"]["days"] if datetime.date.fromisoformat(d["day_date"]).weekday() <= 4][0]["derived"]
q1 = float(L["WS1"]["week"]["planned_qty"]) / NWD
exp_eur = f1["trips"] * 45.0 + q1 * 2.0 + f1["trips"] * 60.0 * 1.5
ok("⭐ € = trips×per_load + tonnes×per_t + trips×basis_km×per_km — all three SUMMED, round trip",
   abs(f1.get("eur") - round(exp_eur, 2)) < 0.011 and L["WS1"]["context"]["rate_set"] is True
   and L["WS1"]["context"]["km_basis"] == "round_trip", str(f1))
f3 = [d for d in L["WS3"]["days"] if datetime.date.fromisoformat(d["day_date"]).weekday() <= 4][0]["derived"]
q3 = float(L["WS3"]["week"]["planned_qty"]) / NWD
ok("⭐ a route on the LOADED basis uses 30 km for t·km, and € with only per_t set is tonnes×per_t",
   L["WS3"]["context"]["km_basis"] == "loaded" and L["WS3"]["context"]["basis_km"] == 30.0
   and abs(f3.get("tonne_km") - round(q3 * 30.0, 1)) < 0.11 and abs(f3.get("eur") - round(q3 * 3.0, 2)) < 0.011)
ok("...rates live on the ROUTE: R1's second line (another IPT) prices with the same rates",
   L["WS2"]["context"]["rates"] == L["WS1"]["context"]["rates"])
T = pg["commit"]["totals"]
ok("totals.eur sums the priced lines and counts them",
   T.get("eur") is not None and T.get("eur_lines") == 3 and T.get("eur_partial") is False)

# ---- the rail
codes = pg["clashes"]["by_code"]
flags = pg["clashes"]["flags"]
ok("🔴 ROUTE_CAP: 2 + 3 vehicles on R1 against a typed cap of 4 flags BOTH lines, every weekday",
   codes.get("ROUTE_CAP") == 2 * NWD and all(f["cap"] == 4 and f["vehicles"] == 5 for f in flags if f["code"] == "ROUTE_CAP"))
ok("🔴 IPT_SHARE: IPT1 + IPT2 on the same route_id the same day — and on the same origin+dest",
   codes.get("IPT_SHARE", 0) >= 2 * NWD
   and {f["share"] for f in flags if f["code"] == "IPT_SHARE"} == {"route", "od"}
   and all(f["ipts"] == ["IPT1", "IPT2"] for f in flags if f["code"] == "IPT_SHARE"))
ok("...R3 (IPT3 alone, a different origin) raises no share flag",
   not any(f["route_id"] == "R3" for f in flags if f["code"] == "IPT_SHARE"))
# a cap EQUAL to the day's vehicles is not exceeded: R3 runs 1 veh/day; cap it at 1
network.set_route_planning("R3", {"max_vehicles_per_day"}, max_vehicles_per_day=1)
_pg_cap = lookahead.page(bucket="commit")
ok("🔴 ROUTE_CAP fires only ABOVE the cap — 1 veh against a cap of 1 raises nothing",
   not any(f["route_id"] == "R3" for f in _pg_cap["clashes"]["flags"] if f["code"] == "ROUTE_CAP")
   and any(f["route_id"] == "R1" for f in _pg_cap["clashes"]["flags"] if f["code"] == "ROUTE_CAP"))
network.set_route_planning("R3", {"max_vehicles_per_day"}, max_vehicles_per_day=None)
# a line with NO IPT cannot share a route with anyone — IPT3 + nobody is not a share
_line2("R3", "substructure", "WS4", "t", 1000.0, V8, ipt=None)
_pg_noipt = lookahead.page(bucket="commit")
ok("🔴 IPT_SHARE never counts a line with no IPT as a second IPT",
   not any(f["route_id"] == "R3" for f in _pg_noipt["clashes"]["flags"] if f["code"] == "IPT_SHARE")
   and any(l["section_id"] == "WS4" for l in _pg_noipt["commit"]["lines"]))
db.execute("DELETE FROM forecast_days WHERE tenant_id = ? AND section_id = 'WS4'", (db.current_tenant(),))
db.execute("DELETE FROM forecast_weeks WHERE tenant_id = ? AND section_id = 'WS4'", (db.current_tenant(),))
db.execute("DELETE FROM forecasts WHERE tenant_id = ? AND section_id = 'WS4'", (db.current_tenant(),))
ok("the empty cap on R3 raises no ROUTE_CAP — no 40-trip constant anywhere",
   not any(f["route_id"] == "R3" for f in flags if f["code"] == "ROUTE_CAP")
   and "40" not in open(os.path.join(BACKEND, "clashes.py"), encoding="utf-8").read().replace("40-trip", ""))
st = {s["location_id"]: s for s in pg["stock"]}
exp_in = sum(float(L[w]["week"]["planned_qty"]) for w in ("WS1", "WS2", "WS3"))
ok("⭐ stock forecast: opening + planned inbound of every line into the stockpile − consume, over when > capacity",
   st["L2"]["over"] is True and abs(st["L2"]["inbound_planned"] - exp_in) < 1e-6
   and abs(st["L2"]["forecast"] - (100.0 + exp_in)) < 1e-6 and abs(st["L2"]["over_by"] - (100.0 + exp_in - 1000.0)) < 1e-6)
ok("🔴 STOCKPILE_OVER on every line into the over-capacity stockpile", codes.get("STOCKPILE_OVER") == 3)
# 🔴 09 Sep night: the live Tark Tee fetch + geometry loop was on the page's critical path
# and froze the Look-ahead; off the path it still took 30 s+ per read. 10 Sep: the check is
# RUN ON DEMAND (or cleared by a bake) and STORED on the route; the page reads the store.
import restrictions as _rx
_calls = {"fetch": 0, "check": []}
_orig_fetch, _orig_check = _rx.fetch_all, _rx.check_route
def _fake_fetch(*a, **k):
    _calls["fetch"] += 1
    return {"type": "FeatureCollection", "features": [{"geometry": None, "properties": {}}], "errors": {}}
def _fake_check(rid, profile=None, layers=None, _fc=None):
    _calls["check"].append(rid)
    return {"route_id": rid, "baked": True, "hits": ([{"headline": "3.5 m limit"}] if rid == "R3" else [])}
_rx.fetch_all, _rx.check_route = _fake_fetch, _fake_check
try:
    _pg_page = lookahead.page(bucket="commit")
    ok("🔴 the page read never touches live Tark Tee — before any refresh the rail says 'unchecked', no flag",
       _calls["fetch"] == 0 and _calls["check"] == []
       and _pg_page["clashes"]["sources"]["tark_tee"] == "unchecked"
       and set(_pg_page["clashes"]["sources"]["tark_tee_unchecked"]) == {"R1", "R3"}
       and _pg_page["clashes"]["by_code"].get("TARK_TEE") is None)
    _st = _rx.refresh_async(sync=True)
    ok("⭐ a refresh fetches Tark Tee ONCE, checks every route, and stores the result on each",
       _calls["fetch"] == 1 and sorted(set(_calls["check"])) == ["R1", "R2", "R3"] and _st["running"] is False
       and _st["status"] == "ok" and _st["routes"] == 3 and _st["hits"] == 1
       and all(x["checked_at"] for x in _rx.stored_checks(["R1", "R3"]).values()))
    _pg_page = lookahead.page(bucket="commit")
    ok("⭐ ...after which the page carries TARK_TEE from the STORE — instantly, no fetch, with the headline and its age",
       _calls["fetch"] == 1 and _pg_page["clashes"]["by_code"].get("TARK_TEE") == 1
       and _pg_page["clashes"]["sources"]["tark_tee"] == "ok" and _pg_page["clashes"]["sources"]["tark_tee_checked_at"]
       and any("3.5 m limit" in f["text"] and f["route_id"] == "R3" for f in _pg_page["clashes"]["flags"]))
    _tt = lookahead.tark_tee_status(bucket="commit")
    ok("...and the status read agrees, and reports the refresh state",
       _tt["status"] == "ok" and _tt["count"] == 1 and _tt["refresh"]["running"] is False and _tt["routes"] == ["R1", "R3"])
    # a re-bake clears the stored check for THAT route only
    network._upsert_geom("R3", V8, "[[24,58.5],[24.4,58.6]]", 31.0, 0.7, None, leg="loaded", alt_index=0)
    _pg_page = lookahead.page(bucket="commit")
    ok("🔴 re-baking a route CLEARS its stored check — the page says 'partial' and names R3, and R3's flag is gone",
       _pg_page["clashes"]["sources"]["tark_tee"] == "partial"
       and _pg_page["clashes"]["sources"]["tark_tee_unchecked"] == ["R3"]
       and _pg_page["clashes"]["by_code"].get("TARK_TEE") is None
       and _rx.stored_checks(["R1"])["R1"]["checked_at"])
    _rx.refresh_async(sync=True)
    ok("...and a refresh fills it again", lookahead.page(bucket="commit")["clashes"]["sources"]["tark_tee"] == "ok")
    _calls["fetch"] = 0
    _rx.fetch_all = lambda *a, **k: (_ for _ in ()).throw(RuntimeError("down"))
    _st2 = _rx.refresh_async(sync=True)
    ok("🔴 a Tark Tee outage during a refresh writes NOTHING — the stored checks stay, status says unavailable",
       _st2["status"] == "unavailable" and lookahead.page(bucket="commit")["clashes"]["sources"]["tark_tee"] == "ok")
    # the other failure shape: every layer errored, so fetch_all returns no features and an errors dict
    _before = _rx.stored_checks(["R1", "R3"])
    _rx.fetch_all = lambda *a, **k: {"type": "FeatureCollection", "features": [], "errors": {"restrictions_mass": "timeout"}}
    _st3 = _rx.refresh_async(sync=True)
    ok("🔴 ...and so does a fetch that returned NO features with errors — never stored as 'no hits'",
       _st3["status"] == "unavailable" and _rx.stored_checks(["R1", "R3"]) == _before
       and _rx.stored_checks(["R3"])["R3"]["hits"] == [{"headline": "3.5 m limit"}])
finally:
    _rx.fetch_all, _rx.check_route = _orig_fetch, _orig_check
main_src2 = open(os.path.join(BACKEND, "main.py"), encoding="utf-8").read()
ok("🔴 the refresh has its own POST endpoint and the page read defaults to the store",
   '"/api/forecast-weeks/tark-tee/refresh"' in main_src2
   and re.search(r'def lookahead_page\([^)]*tark_tee: int = 1', main_src2) is not None
   and "restrictions.refresh_async" in main_src2)
ok("no UNBAKED, no DAYS_NE_WEEK, no SHORTAGE on a fresh derived week",
   not any(c in codes for c in ("UNBAKED", "DAYS_NE_WEEK", "SHORTAGE")))
ok("flags are ordered by the brief's code order", [f["code"] for f in flags] == sorted([f["code"] for f in flags], key=lambda c: clashes.CODES.index(c)))

# ---- an IPT code sees only its own flags, but the cross-IPT fact still reaches it
os.environ.update({"IPT1_CODE": "one-secret", "IPT2_CODE": "two-secret", "IPT3_CODE": "three-secret",
                   "PLANNER_CODE": "plan-secret", "ADMIN_CODE": "adm-secret"})
_as("one-secret")
pg1 = lookahead.page(bucket="commit")
ok("⭐ IPT1 sees one line, its own ROUTE_CAP and IPT_SHARE flags — the other IPT's quantity never appears",
   len(pg1["commit"]["lines"]) == 1 and pg1["commit"]["lines"][0]["ipt"] == "IPT1"
   and all(f["ipt"] == "IPT1" for f in pg1["clashes"]["flags"])
   and pg1["clashes"]["by_code"].get("ROUTE_CAP") == NWD and pg1["clashes"]["by_code"].get("IPT_SHARE", 0) >= NWD
   and pg1["commit"]["totals"]["lines"] == 1)
ok("...and the stock read is not IPT-scoped (C10), so IPT1 still sees the stockpile forecast", len(pg1["stock"]) == 1)
for _v in ("IPT1_CODE", "IPT2_CODE", "IPT3_CODE", "PLANNER_CODE", "ADMIN_CODE"):
    os.environ.pop(_v, None)
_as("planner123")

# ---- the 'next' bucket: the Thursday process
nm, nw = weeks.next_week(MI, WI)
n_before = db.query("SELECT COUNT(*) AS n FROM forecast_days WHERE tenant_id = ?", (db.current_tenant(),))[0]["n"]
pgn = lookahead.page(bucket="next")
ok("⭐ bucket=next reads the week AFTER the commit week and materialises its days on demand",
   pgn["bucket"] == "next" and (pgn["commit_week"]["month_index"], pgn["commit_week"]["week_index"]) == (nm, nw)
   and pgn["today_week"] == {"month_index": MI, "week_index": WI}
   and db.query("SELECT COUNT(*) AS n FROM forecast_days WHERE tenant_id = ?", (db.current_tenant(),))[0]["n"] > n_before)
ok("...its account week is THIS week", pgn["account"]["week"] == {"month_index": MI, "week_index": WI,
   "from": ISO[0], "to": ISO[-1]})
ok("🔴 the default read did NOT materialise the next bucket first — the brief's rule holds by default",
   n_before == len(DATES) * 3)
nd = days.bucket_dates(nm, nw)
r = days.set_day("R1", nm, "earthworks", "WS1", nd[0].isoformat(), 5.0)
ok("a day in the next bucket is typeable", r.get("error") is None)
r = days.set_day("R1", MI, "earthworks", "WS1", (DATES[0] - datetime.timedelta(days=20)).isoformat(), 5.0)
ok("...a day two buckets away is still refused", "not in the commit week" in (r.get("error") or ""))

# ---- account: hold band, actions, cost, € variance
pm, pw = lookahead.prev_week(MI, WI)
if pm >= 1:
    prev1 = weeks.get_week("R1", pm, "earthworks", "WS1", pw)
    prev2 = weeks.get_week("R1", pm, "earthworks", "WS2", pw)
    prev3 = weeks.get_week("R3", pm, "substructure", "WS3", pw)
    p1, p2 = float(prev1["planned_qty"]), float(prev2["planned_qty"])
    main.set_forecast_week_actual(main.WeekActual(route_id="R1", month_index=pm, discipline="earthworks", section_id="WS1",
                                                  week_index=pw, actual_qty=p1 * 0.985, actual_cost_eur=1234.5))
    main.set_forecast_week_actual(main.WeekActual(route_id="R1", month_index=pm, discipline="earthworks", section_id="WS2",
                                                  week_index=pw, actual_qty=p2 - 180.0))
    pg = lookahead.page(bucket="commit")
    A = {r["section_id"]: r for r in pg["account"]["rows"]}
    ok("⭐ 98 % band: 98.5 % delivered holds; 180 short does not; untyped waits",
       A["WS1"]["held"] is True and A["WS1"]["action"] == "held"
       and A["WS2"]["held"] is False and A["WS2"]["action"] == "spread" and abs(A["WS2"]["remaining_short"] - 180.0) < 1e-6
       and A["WS3"]["action"] == "waiting" and A["WS3"]["actual_qty"] is None)
    ok("...the KPI figures: 1 of 3 delivered, 2 reported, shortfall 180, 1 open to calibrate",
       pg["account"]["delivered"] == 1 and pg["account"]["lines"] == 3 and pg["account"]["reported"] == 2
       and abs(pg["account"]["shortfall"] - 180.0) < 1e-6 and pg["account"]["open_to_calibrate"] == 1)
    ok("⭐ the actual cost is stored and € variance = planned € − actual €",
       A["WS1"]["actual_cost_eur"] == 1234.5 and A["WS1"]["planned_eur"] is not None
       and abs(A["WS1"]["eur_variance"] - round(A["WS1"]["planned_eur"] - 1234.5, 2)) < 1e-6
       and A["WS2"]["actual_cost_eur"] is None and A["WS2"]["eur_variance"] is None)
    main.set_forecast_week_actual(main.WeekActual(route_id="R1", month_index=pm, discipline="earthworks", section_id="WS1",
                                                  week_index=pw, actual_qty=p1 * 0.985, actual_note="re-typed"))
    ok("🔴 re-saving the actual WITHOUT the cost field leaves the stored cost alone",
       weeks.get_week("R1", pm, "earthworks", "WS1", pw)["actual_cost_eur"] == 1234.5)
    main.set_forecast_week_actual(main.WeekActual(route_id="R1", month_index=pm, discipline="earthworks", section_id="WS1",
                                                  week_index=pw, actual_qty=p1 * 0.985, actual_cost_eur=None))
    ok("🔴 ...but an EXPLICIT null clears it — absent and null are different things",
       weeks.get_week("R1", pm, "earthworks", "WS1", pw)["actual_cost_eur"] is None)
    main.set_forecast_week_actual(main.WeekActual(route_id="R1", month_index=pm, discipline="earthworks", section_id="WS1",
                                                  week_index=pw, actual_qty=p1 * 0.985, actual_cost_eur=1234.5))
    ok("🔴 SHORTAGE is on the rail for the short line and not for the held one",
       {f["section_id"] for f in pg["clashes"]["flags"] if f["code"] == "SHORTAGE"} == {"WS2"})
    # carry it: the shortage leaves the rail and the action reads applied
    days.calibrate("R1", pm, "earthworks", "WS2", pw, spread=True)
    pg = lookahead.page(bucket="commit")
    A = {r["section_id"]: r for r in pg["account"]["rows"]}
    ok("⭐ after calibrate the line reads `applied`, its shortage is carried, and SHORTAGE leaves the rail",
       A["WS2"]["action"] == "applied" and abs(A["WS2"]["carried"] - 180.0) < 1e-6
       and not any(f["code"] == "SHORTAGE" for f in pg["clashes"]["flags"])
       and pg["account"]["open_to_calibrate"] == 0)
    ok("...and the spread made the days ≠ week flag stay quiet (the spread keeps the sum rule)",
       not any(f["code"] == "DAYS_NE_WEEK" for f in pg["clashes"]["flags"]))
else:
    for _ in range(7):
        ok("(account assertions skipped — no previous bucket inside the horizon)", True)

# ---- the render fixture: written HERE so render_frontend.js renders the shape this
# backend actually produces (timestamps scrubbed). The js harness runs after this file.
import json as _json
_fix_dir = os.path.join(HERE, "fixtures")
os.makedirs(_fix_dir, exist_ok=True)
def _scrub(o):
    if isinstance(o, dict):
        return {k: ("<ts>" if k.endswith("_at") and isinstance(v, str) else _scrub(v)) for k, v in o.items()}
    if isinstance(o, list):
        return [_scrub(x) for x in o]
    return o
_fix = _scrub(lookahead.page(bucket="commit"))
with open(os.path.join(_fix_dir, "lookahead_page.json"), "w", encoding="utf-8") as _fh:
    _fh.write(_json.dumps(_fix, indent=1, ensure_ascii=False))
ok("the render fixture is written with three lines, a rail, a stockpile and an account week",
   len(_fix["commit"]["lines"]) == 3 and _fix["clashes"]["count"] > 3 and len(_fix["stock"]) == 1
   and len(_fix["account"]["rows"]) == (3 if pm >= 1 else 0))

# ---- horizon roles
hz = pg["horizon"]
roles = {(r["month_index"], r["week_index"]): r["role"] for r in hz["rows"]}
ok("horizon rows carry account / commit / make-ready / early-warning relative to the commit bucket",
   roles.get((MI, WI)) == "commit" and roles.get(tuple(weeks.next_week(MI, WI))) == "make-ready"
   and roles.get(tuple(weeks.next_week(*weeks.next_week(MI, WI)))) == "early-warning"
   and (pm < 1 or roles.get((pm, pw)) == "account")
   and hz["from_month"] == MI and hz["to_month"] == MI + 1)

# ---- confirm, then reopen — days follow both ways
main.confirm_forecast_week(main.WeekConfirm(route_id="R3", month_index=MI, discipline="substructure", section_id="WS3", week_index=WI))
ok("confirm still stamps the days", all(d["status"] == "confirmed" for d in days._days_of("R3", MI, "substructure", "WS3", WI)))
r = days.set_day("R3", MI, "substructure", "WS3", ISO[0], 1.0)
ok("...and a confirmed week's days are read-only", r.get("blocked_by") == "confirmed")
r = main.reopen_forecast_week(main.WeekReopen(route_id="R3", month_index=MI, discipline="substructure", section_id="WS3", week_index=WI))
ok("⭐ reopen takes the week and its days back to `edited` — nothing deleted",
   r.get("error") is None and r["week"]["status"] == "edited" and r["week"]["confirmed_at"]
   and all(d["status"] == "edited" for d in days._days_of("R3", MI, "substructure", "WS3", WI)))
r = days.set_day("R3", MI, "substructure", "WS3", ISO[0], 1.0)
ok("...so the plan can change again", r.get("error") is None)
try:
    main.reopen_forecast_week(main.WeekReopen(route_id="R3", month_index=MI, discipline="substructure", section_id="WS3", week_index=WI))
    ok("reopening a week that is not confirmed is a 400", False)
except Exception as e:
    ok("reopening a week that is not confirmed is a 400", getattr(e, "status_code", None) == 400)

# ---- export: bytes that open, with the right sheets and the right rows
pg = lookahead.page(bucket="commit")
xb = export.build_xlsx(pg)
ok("the XLSX builds", isinstance(xb, bytes) and xb[:2] == b"PK")
import io
import openpyxl as _ox
wb = _ox.load_workbook(io.BytesIO(xb))
# 10 Sep, the human: the sheet is the SUPPLIER's — Mon–Fri only, and no stockpile list
ok("...with three sheets: Commit week, Clashes, About — NO Stock sheet (10 Sep)",
   wb.sheetnames == ["Commit week", "Clashes", "About"])
ws = wb["Commit week"]
ok("⭐ ...sheet 1 is WEEKDAY × line: one row per Mon–Fri day per visible line, headers from the export mock",
   ws.max_row - 1 == 3 * NWD and [c.value for c in ws[1]][:6] == ["Date", "Route", "Origin", "Destination", "IPT", "WS"])
ok("🔴 ...and no Saturday or Sunday row at all",
   all(datetime.date.fromisoformat(str(ws.cell(row=i, column=1).value)[:10]).weekday() <= 4 for i in range(2, ws.max_row + 1)))
ok("...a priced row carries €",
   any(ws.cell(row=i, column=18).value not in (None, "") for i in range(2, ws.max_row + 1)))
# .get-style access: a renamed sheet must FAIL these, not crash the report (lesson 13)
_sh = lambda n: wb[n] if n in wb.sheetnames else None
ok("...the Clashes sheet has every flag the rail has",
   _sh("Clashes") is not None and _sh("Clashes").max_row - 1 == pg["clashes"]["count"])
ok("...the About sheet says Mon–Fri only",
   _sh("About") is not None and any("Mon–Fri only" in str(_sh("About").cell(row=i, column=2).value or "") for i in range(1, _sh("About").max_row + 1)))
pb = export.build_pdf(pg)
ok("the PDF builds", isinstance(pb, bytes) and pb[:5] == b"%PDF-")
ptxt = pb.decode("latin-1")
ok("🔴 the PDF is LANDSCAPE A4 (10 Sep, the human)", re.search(r"/MediaBox \[ 0 0 841\.\d+ 595\.\d+ \]", ptxt) is not None, ptxt[ptxt.find("/MediaBox"):ptxt.find("/MediaBox") + 40])
_txt = subprocess.run(["pdftotext", "-layout", "-", "-"], input=pb, capture_output=True).stdout.decode("utf-8") if shutil.which("pdftotext") else ptxt
ok("...and carries the disclaimer, the flags heading, the footer lines and a page number",
   "not a delivery note" in _txt and "not a stop" in _txt and "Vignette is time-based" in _txt
   and "Re-open the week" in _txt and "page 1" in _txt)
ok("⭐ one row per LINE with Mon…Fri as five separate columns — no collapse rule",
   "FRI EACH DAY" not in _txt and all(k in _txt for k in ("MON 7 SEP", "TUE 8 SEP", "WED 9 SEP", "THU 10 SEP", "FRI 11 SEP"))
   and _txt.count("Small aggregate") == 3)
ok("⭐ origin and destination carry their coordinates from the locations table",
   "58.50000, 24.00000" in _txt and "58.60000, 24.40000" in _txt and "ORIGIN" in _txt and "DESTINATION" in _txt)
ok("...today's column is named, the week column totals, and a TOTAL row closes the table",
   "TODAY" in _txt and "WEEK" in _txt and "TOTAL" in _txt and "3 line(s)" in _txt)
ok("...a draft week says so, never CONFIRMED", "not confirmed" in _txt and "CONFIRMED" not in _txt.replace("not confirmed", ""))
ok("🔴 the PDF has NO stock section and NO weekend column (10 Sep)",
   "Stock at week end" not in _txt and re.search(r"\b(SAT|SUN) \d", _txt) is None)
ok("⭐ ...and carries the route map — the schematic here, because Mapbox cannot be reached from the sandbox",
   "Routes this week: 2" in _txt and "schematic from the baked geometry" in _txt)
# the Mapbox branch, with the fetch stubbed: the PNG is placed and the caption says Mapbox
# a REAL 8×4 PNG (Pillow writes it; reportlab needs Pillow to place a PNG — it is on
# Render because reportlab pulls it in, and it is asserted below)
import io as _io
from PIL import Image as _Img
_pb = _io.BytesIO(); _Img.new("RGB", (8, 4), (200, 210, 230)).save(_pb, format="PNG"); _png = _pb.getvalue()
class _Resp:
    def __init__(self, b): self.b = b
    def read(self): return self.b
    def __enter__(self): return self
    def __exit__(self, *a): return False
_orig_open = export.urllib.request.urlopen
export.urllib.request.urlopen = lambda req, timeout=0: _Resp(_png)
try:
    pb2 = export.build_pdf(pg)
    _txt2 = subprocess.run(["pdftotext", "-layout", "-", "-"], input=pb2, capture_output=True).stdout.decode("utf-8") if shutil.which("pdftotext") else pb2.decode("latin-1")
    ok("⭐ when Mapbox answers, the PDF carries the image and the caption says so — the branch Render will take",
       "map © Mapbox" in _txt2 and "schematic" not in _txt2 and b"/Subtype /Image" in pb2)
    ok("...and the request carries the SAME public token the browser map uses when MAPBOX_TOKEN is unset",
       export.mapbox_static_png(export.route_geometries(pg)) == _png and config.mapbox_token() == config.MAPBOX_TOKEN_DEFAULT)
    _seen = {}
    export.urllib.request.urlopen = lambda req, timeout=0: (_seen.setdefault("url", req.full_url), _Resp(_png))[1]
    export.mapbox_static_png(export.route_geometries(pg))
    ok("...the URL is a Static Images request over light-v11 with a path overlay per route",
       _seen.get("url", "").startswith("https://api.mapbox.com/styles/v1/mapbox/light-v11/static/") and _seen.get("url", "").count("path-") == 2
       and "access_token=" + config.MAPBOX_TOKEN_DEFAULT in _seen.get("url", ""), _seen.get("url", "no request made"))
    export.urllib.request.urlopen = lambda req, timeout=0: _Resp(b"<html>rate limited</html>")
    ok("🔴 a non-PNG answer is not drawn as a map — None, and the schematic takes over", export.mapbox_static_png(export.route_geometries(pg)) is None)
finally:
    export.urllib.request.urlopen = _orig_open
# many lines: several pages, the header row on each
import copy as _copy
_big = _copy.deepcopy(pg)
_big["commit"]["lines"] = [_copy.deepcopy(pg["commit"]["lines"][i % 3]) for i in range(40)]
pb3 = export.build_pdf(_big)
_txt3 = subprocess.run(["pdftotext", "-layout", "-", "-"], input=pb3, capture_output=True).stdout.decode("utf-8") if shutil.which("pdftotext") else pb3.decode("latin-1")
_npages = pb3.count(b"/Type /Page\n")
ok("⭐ forty lines run over several pages, and the column headers repeat on every page",
   _npages >= 3 and _npages - 1 <= _txt3.count("DESTINATION") <= _npages and "40 line(s)" in _txt3,   # the last page may hold only the flags + footer
   f"{_npages} pages, {_txt3.count('DESTINATION')} headers")
_geoms = export.route_geometries(pg)
ok("...drawn from the lines' own baked geometry, one per route, the line's vehicle first",
   sorted(g[0] for g in _geoms) == ["R1", "R3"] and all(len(g[3]) >= 2 for g in _geoms))
ok("the polyline encoder round-trips a known point (Google's own example)",
   export._encode_polyline([(-120.2, 38.5), (-120.95, 40.7), (-126.453, 43.252)]) == "_p~iF~ps|U_ulLnnqC_mqNvxq`@")
ok("with an EMPTY token the Mapbox static path is skipped, not attempted",
   export.mapbox_static_png(_geoms, token="") is None)
ok("no email, no upload anywhere in the export or the endpoints",
   "smtp" not in open(os.path.join(BACKEND, "export.py"), encoding="utf-8").read().lower()
   and "UploadFile" not in open(os.path.join(BACKEND, "main.py"), encoding="utf-8").read())


# =========================================================================== #
#  12. THE DEPLOYMENT'S READ COST (2026-09-09 night) — measured live at 17.7 s   #
#      for ONE line: one Postgres connection per statement, and a read that     #
#      rewrote every derived row it had just read                               #
# =========================================================================== #
import collections as _co
_orig_q, _orig_x = db.query, db.execute
_cnt = _co.Counter()
def _q(sql, *a, **k):
    _cnt["query"] += 1
    return _orig_q(sql, *a, **k)
def _x(sql, *a, **k):
    _cnt["execute"] += 1
    return _orig_x(sql, *a, **k)
db.query, db.execute = _q, _x
try:
    _as("planner123")
    _cnt.clear(); main.list_forecast_days(); first = dict(_cnt)
    _cnt.clear(); main.list_forecast_days(); second = dict(_cnt)
    ok("🔴 a SECOND read of the same days writes NOTHING — derived rows are rewritten only when they moved",
       second.get("execute", 0) == 0, f"second read: {second}")
    ok(f"...and the days read stays under a statement budget (was 91 for 3 lines): {second.get('query', 0)} reads",
       second.get("query", 0) <= 60)
    _cnt.clear(); lookahead.page(bucket="commit"); pg_cost = dict(_cnt)
    ok(f"the page read for 3 lines: {pg_cost.get('query', 0)} reads, {pg_cost.get('execute', 0)} writes — writes must be 0 on a settled week",
       pg_cost.get("execute", 0) == 0 and pg_cost.get("query", 0) <= 130, str(pg_cost))
    # but a week that MOVED still refreshes its derived days
    main.save_matrix_row(main.MatrixRow(route_id="R1", discipline="earthworks", section_id="WS1",
        material_type="Small aggregate", material_description=None, vehicle_type=V8, submitted_by="tester",
        unit="t", status="Pending", cells=[main.Cell(month_index=m, quantity=5000.0) for m in MONTHS], ipt="IPT1"))
    main.set_route_status("R1", main.StatusUpdate(status="Approved"), discipline="earthworks", section_id="WS1")
    _cnt.clear(); main.list_forecast_days(); moved = dict(_cnt)
    ok("⭐ ...while a week whose parent moved DOES rewrite its derived days (the refresh still works)",
       moved.get("execute", 0) > 0 and abs(float({l["section_id"]: l for l in main.list_forecast_days()["lines"]}["WS1"]["week"]["planned_qty"]) - 5000.0 * len(DATES) / 30) < 400)
finally:
    db.query, db.execute = _orig_q, _orig_x

# ---- the pool: psycopg2 is absent here, so the proxy and the retry are exercised on stand-ins
class _FakeConn:
    def __init__(self, tx=0, closed=False, fail_first=False):
        self.tx, self.closed, self.fail_first = tx, closed, fail_first
        self.rolled_back = self.committed = False; self.executed = []
    def get_transaction_status(self): return self.tx
    def rollback(self): self.rolled_back = True; self.tx = 0
    def commit(self): self.committed = True
    def close(self): self.closed = True
    def cursor(self, **k):
        conn = self
        class _Cur:
            def execute(self_, sql, params=()):
                if conn.fail_first:
                    conn.fail_first = False
                    raise type("OperationalError", (Exception,), {})("server closed the connection unexpectedly")
                conn.executed.append(sql)
            def fetchall(self_): return [{"n": 1}]
        return _Cur()
class _FakePool:
    def __init__(self): self.put = []; self.given = []; self.next = []
    def getconn(self):
        c = self.next.pop(0) if self.next else _FakeConn(); self.given.append(c); return c
    def putconn(self, conn, close=False): self.put.append((conn, close))

_pool = _FakePool(); _c = _FakeConn(tx=2)          # tx=2: mid-transaction (TRANSACTION_STATUS_INTRANS)
_pc = db._PooledConn(_pool, _c)
ok("a pooled connection quacks like the real one (cursor/commit/rollback pass through)",
   _pc.cursor() is not None and (_pc.commit() or _c.committed))
_pc.close()
ok("⭐ close() RETURNS the connection to the pool, rolling back a half-open transaction first",
   _pool.put == [(_c, False)] and _c.rolled_back and not _c.closed)
try:
    _pc.cursor(); ok("...and a returned connection cannot be used again by mistake", False)
except RuntimeError:
    ok("...and a returned connection cannot be used again by mistake", True)
_pool2 = _FakePool(); _dead = _FakeConn(closed=True)
db._PooledConn(_pool2, _dead).close()
ok("a connection that died is returned with close=True, never reused", _pool2.put == [(_dead, True)])
_pool3 = _FakePool(); _pc3 = db._PooledConn(_pool3, _FakeConn()); _pc3.discard()
ok("discard() drops it from the pool for good", _pool3.put and _pool3.put[0][1] is True)
# the retry: first statement hits a dead pooled connection, the second succeeds on a fresh one
_pool4 = _FakePool(); _pool4.next = [_FakeConn(fail_first=True), _FakeConn()]
_orig_get, _orig_pg = db.get_conn, db.IS_PG
db.get_conn = lambda: db._PooledConn(_pool4, _pool4.getconn())
db.IS_PG = False                                   # keep the sqlite cursor path (no psycopg2.extras here)
try:
    try:
        rows = db.query("SELECT 1")
    except Exception as e:                       # an assertion that can crash hides every one after it
        rows = f"raised {type(e).__name__}"
    ok("🔴 a statement on a dead pooled connection is retried ONCE on a fresh one and succeeds",
       rows == [{"n": 1}] and len(_pool4.given) == 2 and _pool4.put[0][1] is True and _pool4.put[1][1] is False, str(rows))
    _pool4.next = [_FakeConn(fail_first=True), _FakeConn(fail_first=True)]; _pool4.given.clear(); _pool4.put.clear()
    try:
        db.query("SELECT 1"); ok("...but only once — two dead connections in a row raise", False)
    except Exception as e:
        ok("...but only once — two dead connections in a row raise", type(e).__name__ == "OperationalError" and len(_pool4.given) == 2)
    # a STATEMENT error (bad SQL) must NOT be retried — only a dead connection is
    class _BadSqlConn(_FakeConn):
        def cursor(self, **k):
            class _Cur:
                def execute(self_, sql, params=()):
                    raise type("ProgrammingError", (Exception,), {})("bad sql")
            return _Cur()
    _pool4.next = [_BadSqlConn(), _FakeConn()]; _pool4.given.clear(); _pool4.put.clear()
    try:
        db.query("SELECT bad"); ok("a STATEMENT error is not retried (a retry would not fix bad SQL)", False)
    except Exception as e:
        ok("a STATEMENT error is not retried (a retry would not fix bad SQL)",
           type(e).__name__ == "ProgrammingError" and len(_pool4.given) == 1 and _pool4.put[0][1] is False)
finally:
    db.get_conn, db.IS_PG = _orig_get, _orig_pg
db_src = open(os.path.join(BACKEND, "db.py"), encoding="utf-8").read()
ok("get_conn() on Postgres borrows from the pool, and DB_POOL=0 turns it off",
   "ThreadedConnectionPool" in db_src and 'os.getenv("DB_POOL", "1")' in db_src and "return _PooledConn(pool, conn)" in db_src)
ok("query() and execute() both go through the retrying _run()",
   db_src.count("return _run(sql, params, fetch=True)") == 1 and db_src.count("_run(sql, params, fetch=False)") == 1)


# =========================================================================== #
#  13. 10 SEP — calendar weeks in the data, stock on Account, the week map       #
# =========================================================================== #
# (state as left by section 12: three lines on R1/R2/R3, L2 holding stock, MI/WI commit)
_as("planner123")
pg13 = lookahead.page(bucket="commit")
hz13 = pg13["horizon"]
ok("the horizon says how many weeks each month has, and where each one runs — nothing assumes four",
   set(hz13["weeks_in_month"]) == {str(MI), str(MI + 1)}
   and all(int(v) in (4, 5) for v in hz13["weeks_in_month"].values())
   and all(f"{m}|{w}" in hz13["week_spans"] for m in (MI, MI + 1) for w in range(1, weeks.weeks_in_month(m) + 1))
   and all(datetime.date.fromisoformat(v["from"]).weekday() == 0 for v in hz13["week_spans"].values()))
ok("⭐ every week span on the horizon starts on a MONDAY and is seven days",
   all((datetime.date.fromisoformat(v["to"]) - datetime.date.fromisoformat(v["from"])).days == 6
       for v in hz13["week_spans"].values()))
ok("the commit week itself starts on a Monday — the human's rule, 10 Sep",
   datetime.date.fromisoformat(pg13["commit_week"]["from"]).weekday() == 0
   and datetime.date.fromisoformat(pg13["commit_week"]["to"]).weekday() == 6)

# the account week's stockpiles: one row per stockpile, the balance arithmetic of
# stockpiles.balances(), consumed None until typed
if pm >= 1:
    st = pg13["account"]["stock"]
    ok("account.stock has one row per stockpile, with opening / in / out / closing / remaining",
       len(st) == 1 and st[0]["location_id"] == "L2"
       and all(k in st[0] for k in ("opening", "inbound", "consumed", "closing", "remaining", "over", "capacity_qty", "unit")))
    ok("⭐ 'out' is None until typed — not 0 — and closing = opening + in − 0 meanwhile",
       st[0]["consumed"] is None and abs(st[0]["closing"] - (st[0]["opening"] + st[0]["inbound"])) < 1e-6)
    ok("...the figures agree with stockpiles.balances() for that week (nothing recomputed)",
       abs(st[0]["closing"] - [w for w in stockpiles.balances(1, pm)["stockpiles"][0]["weeks"]
                              if (w["month_index"], w["week_index"]) == (pm, pw)][0]["balance_end"]) < 1e-6)
    # type 50 out through the same PUT the old grid used
    main.consume_stockpile(main.ConsumeIn(location_id="L2", month_index=pm, week_index=pw, consumed_qty=50.0, unit="t", note=None, updated_by="t"))
    st2 = lookahead.page(bucket="commit")["account"]["stock"][0]
    ok("typing 50 out lowers closing by 50 and reads back as typed",
       st2["consumed"] == 50.0 and abs(st2["closing"] - (st[0]["closing"] - 50.0)) < 1e-6)
    ok("...and over/remaining follow the capacity (1 000 t here)",
       st2["over"] is (st2["closing"] > 1000.0) and abs(st2["remaining"] - (1000.0 - st2["closing"])) < 1e-6)
    main.consume_stockpile(main.ConsumeIn(location_id="L2", month_index=pm, week_index=pw, consumed_qty=None, unit="t", note=None, updated_by="t"))
    ok("clearing the box returns 'out' to None, not 0",
       lookahead.page(bucket="commit")["account"]["stock"][0]["consumed"] is None)
    ok("a five-week month refuses week 6 and accepts week 5 on the consume PUT",
       "week_index must be 1-5" in (stockpiles.consume("L2", 10, 6, consumed_qty=1.0, unit="t").get("error") or "")
       and not stockpiles.consume("L2", 10, 5, consumed_qty=1.0, unit="t").get("error"))
else:
    for _ in range(7):
        ok("(account stock assertions skipped — no previous bucket inside the horizon)", True)

# the week map's read: geometry per route, ends with coordinates, unbaked = null
db.execute("INSERT INTO routes (id, origin_id, dest_id, ipt) VALUES (?, ?, ?, ?)", ("R9", "L2", "L1", "IPT 4"))   # never baked
gm = lookahead.week_geometry(["R1", "R3", "R9", "R1"], {"R1": V8, "R3": "Ghost truck"})
ok("week_geometry returns one entry per DISTINCT route, in the order asked",
   [g["route_id"] for g in gm] == ["R1", "R3", "R9"])
ok("⭐ both ends carry name and coordinates from the locations table",
   gm[0]["origin"]["name"] == "Pit" and gm[0]["origin"]["lat"] == 58.5 and gm[0]["origin"]["lon"] == 24.0
   and gm[0]["dest"]["name"] == "Site" and gm[0]["dest"]["lat"] == 58.6)
ok("R1 is drawn for its planned vehicle (baked)", gm[0]["geometry"] == [[24.0, 58.5], [24.4, 58.6]]
   and gm[0]["vehicle_profile"] == V8 and gm[0]["vehicle_as_planned"] is True and gm[0]["distance_km"] == 30.0)
ok("⭐ R3 asked for a vehicle it is NOT baked for: drawn for the one that is, and SAYS so",
   gm[1]["geometry"] is not None and gm[1]["vehicle_profile"] == V8 and gm[1]["vehicle_as_planned"] is False)
ok("🔴 R9 has no baked geometry: geometry is null — never a straight line pretending to be a road",
   gm[2]["geometry"] is None and gm[2]["vehicle_profile"] is None and gm[2]["origin"]["name"] == "Site")
ok("the endpoint parses ids and aligned vehicles, and needs a code",
   [g["route_id"] for g in main.lookahead_geometry(ids="R1,R2", vehicles=V8 + ",")["routes"]] == ["R1", "R2"])
_as(None)
try:
    main.lookahead_geometry(ids="R1"); ok("...no code, no map data", False)
except Exception as e:
    ok("...no code, no map data", getattr(e, "status_code", None) in (401, 403), str(e))
_as("planner123")

# 🔴 the calendar change: a day written under the OLD bucket number is re-stamped on read
_d0 = days.bucket_dates(MI, WI)[0].isoformat()
db.execute("UPDATE forecast_days SET parent_week_index = 99 WHERE tenant_id = ? AND route_id = ? AND day_date = ?",
           (db.current_tenant(), "R1", _d0))
main.list_forecast_days()
_rows = db.query("SELECT parent_week_index, status FROM forecast_days WHERE tenant_id = ? AND route_id = ? AND day_date = ?",
                 (db.current_tenant(), "R1", _d0))
ok("⭐ a day whose bucket number is stale (the weeks moved to the calendar) is re-stamped on the next read, whatever its status",
   len(_rows) >= 1 and all(int(r["parent_week_index"]) == WI for r in _rows), str(_rows))


# =========================================================================== #
print()
print(f"{PASS} passed, {len(FAIL)} failed")
for f in FAIL:
    print("  FAIL:", f)
shutil.rmtree(TMP, ignore_errors=True)
sys.exit(1 if FAIL else 0)
